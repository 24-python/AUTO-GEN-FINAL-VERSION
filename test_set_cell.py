# test_set_cell.py
import sys
from pathlib import Path

sys.path.insert(0, '.')

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import copy
from docx.oxml import OxmlElement

doc = Document("tech_card_templates/шаблон.docx")
table = doc.tables[1]

source_row = table.rows[7]
new_row = table.add_row()

# Клонируем + bold для колонки 4
for col_idx, source_cell in enumerate(source_row.cells):
    target_cell = new_row.cells[col_idx]
    source_tc_pr = source_cell._tc.find(qn('w:tcPr'))
    if source_tc_pr is not None:
        target_tc_pr = target_cell._tc.find(qn('w:tcPr'))
        if target_tc_pr is not None:
            target_cell._tc.remove(target_tc_pr)
        target_cell._tc.insert(0, copy.deepcopy(source_tc_pr))

    for p_idx, source_para in enumerate(source_cell.paragraphs):
        if p_idx >= len(target_cell.paragraphs):
            target_para = target_cell.add_paragraph()
        else:
            target_para = target_cell.paragraphs[p_idx]

        source_p_pr = source_para._p.find(qn('w:pPr'))
        if source_p_pr is not None:
            target_p_pr = target_para._p.find(qn('w:pPr'))
            if target_p_pr is not None:
                target_para._p.remove(target_p_pr)
            target_para._p.insert(0, copy.deepcopy(source_p_pr))

        for run in target_para.runs:
            target_para._p.remove(run._r)

        for source_run in source_para.runs:
            new_run = copy.deepcopy(source_run._r)
            r_pr = new_run.find(qn('w:rPr'))
            if r_pr is None:
                r_pr = OxmlElement('w:rPr')
                new_run.insert(0, r_pr)

            if col_idx == 3:
                b = OxmlElement('w:b')
                r_pr.append(b)

            sz = r_pr.find(qn('w:sz'))
            if sz is None:
                sz = OxmlElement('w:sz')
                r_pr.append(sz)
            sz.set(qn('w:val'), '14')

            r_fonts = r_pr.find(qn('w:rFonts'))
            if r_fonts is None:
                r_fonts = OxmlElement('w:rFonts')
                r_pr.append(r_fonts)
            r_fonts.set(qn('w:ascii'), 'Arial')
            r_fonts.set(qn('w:hAnsi'), 'Arial')

            for t in new_run.findall(qn('w:t')):
                t.text = ''
                break

            target_para._p.append(new_run)

# Проверяем bold ДО _set_cell_text
cell = new_row.cells[3]
run_before = cell.paragraphs[0].runs[0]
r_pr_before = run_before._r.find(qn('w:rPr'))
bold_before = r_pr_before.find(qn('w:b')) if r_pr_before is not None else None
print(f"ДО: bold={'✅' if bold_before is not None else '❌'}")

# Эмулируем _set_cell_text
para = cell.paragraphs[0]
run = para.runs[0]
for t in run._r.findall(qn('w:t')):
    t.text = 'ХИМИТЕК ИЗУМРУД 310'
    break

# Проверяем bold ПОСЛЕ
run_after = cell.paragraphs[0].runs[0]
r_pr_after = run_after._r.find(qn('w:rPr'))
bold_after = r_pr_after.find(qn('w:b')) if r_pr_after is not None else None
print(f"ПОСЛЕ: text='{run_after.text}', bold={'✅' if bold_after is not None else '❌'}")

doc.save("test_set_cell_output.docx")
print("\n✅ Сохранено: test_set_cell_output.docx")