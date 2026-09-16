# generator/docx_generator.py

import traceback
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pathlib import Path
from parser.models import ChecklistData, ChecklistItem, Category
from db.models import (Instruction, Category as DBCategory, Object as DBObject,
                       RoomCategory, Product, ObjectProperty, ObjectGroup,
                       CleaningMethodOrder, InventoryColor, CleaningTechnique)
from db.database import SessionLocal
from collections import defaultdict
import re
import copy


class TechCardGenerator:
    TEMPLATES_DIR = Path("tech_card_templates")
    DEFAULT_TEMPLATE = TEMPLATES_DIR / "шаблон.docx"

    LEVEL_ORDER = {"основная": 1, "поддерживающая": 2, "генеральная": 3}

    SKIP_PLACEHOLDERS = {
        "Выберите элемент.", "Не выбрано", "Выберите элемент",
        "", None
    }

    def __init__(self, template_path: str = None):
        self.template_path = Path(template_path) if template_path else self.DEFAULT_TEMPLATE
        self._config_loaded = False
        self._inventory_colors = {}
        self._cleaning_techniques = {}
        self._load_inventory_colors()
        self._load_cleaning_techniques()

    def _load_inventory_colors(self):
        session = SessionLocal()
        try:
            colors = session.query(InventoryColor).all()
            self._inventory_colors = {c.name: c.hex_color for c in colors}
        finally:
            session.close()

    def _load_cleaning_techniques(self):
        session = SessionLocal()
        try:
            techniques = session.query(CleaningTechnique).all()
            self._cleaning_techniques = {t.name.lower(): t.name for t in techniques}
        finally:
            session.close()

    def _load_config(self, session):
        if self._config_loaded:
            return
        methods = session.query(CleaningMethodOrder).order_by(CleaningMethodOrder.sort_order).all()
        self.cleaning_method_order = {m.method_name: m.sort_order for m in methods}

        props = session.query(ObjectProperty).join(DBObject).all()
        self.split_objects = set()
        self.multi_method_objects = set()
        self.support_objects = set()
        self.floor_objects = set()
        self.tech_objects = set()
        self.glass_objects = set()
        self.thermal_objects = set()

        for prop in props:
            name = prop.object.normalized_name
            if prop.is_split:
                self.split_objects.add(name)
            if prop.is_multi_method:
                self.multi_method_objects.add(name)
            if prop.has_support_maintenance:
                self.support_objects.add(name)
            if prop.special_product_type == 'floor':
                self.floor_objects.add(name)
            elif prop.special_product_type == 'tech':
                self.tech_objects.add(name)
            elif prop.special_product_type == 'glass':
                self.glass_objects.add(name)
            elif prop.special_product_type == 'thermal':
                self.thermal_objects.add(name)

        groups = session.query(ObjectGroup).all()
        self.group_headers = {}
        temp_groups = defaultdict(set)
        for g in groups:
            if g.object:
                temp_groups[g.group_name].add(g.object.normalized_name)
        for header, obj_set in temp_groups.items():
            self.group_headers[frozenset(obj_set)] = header

        self._config_loaded = True

    def _load_product_colors(self, session) -> dict:
        products = session.query(Product).all()
        return {p.name: p.color for p in products if p.color}

    def _clean_text(self, val: str) -> str:
        if not val:
            return ""
        return val.strip().replace('\n', ' ').replace('\r', ' ')

    def _normalize_product_name(self, name: str) -> str:
        if not name:
            return ""
        normalized = re.sub(r'[^\w\s\-]', '', name)
        normalized = normalized.lower()
        normalized = re.sub(r'\s+', ' ', normalized)
        return normalized.strip()

    def _set_cell_background(self, cell, hex_color: str):
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), hex_color)
        cell._tc.get_or_add_tcPr().append(shading)

    def _merge_cells_horizontal(self, row, start_col: int, end_col: int):
        if start_col >= end_col:
            return
        start_cell = row.cells[start_col]
        for col in range(start_col + 1, end_col + 1):
            start_cell.merge(row.cells[col])

    def _remove_extra_paragraphs_in_cell(self, cell):
        tc = cell._tc
        paragraphs = tc.findall(qn('w:p'))
        if not paragraphs:
            return
        first_p = paragraphs[0]
        for p in paragraphs[1:]:
            tc.remove(p)
        if not first_p.findall(qn('w:r')):
            run = OxmlElement('w:r')
            run_pr = OxmlElement('w:rPr')
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '14')
            run_pr.append(sz)
            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), 'Arial')
            rFonts.set(qn('w:hAnsi'), 'Arial')
            run_pr.append(rFonts)
            run.append(run_pr)
            t = OxmlElement('w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = ''
            run.append(t)
            first_p.append(run)

    def _merge_cells_vertical(self, table, col: int, start_row: int, end_row: int):
        if start_row >= end_row:
            return
        start_cell = table.cell(start_row, col)
        for row in range(start_row + 1, end_row + 1):
            start_cell.merge(table.cell(row, col))
        self._remove_extra_paragraphs_in_cell(start_cell)

    def _cells_are_equal(self, table, col: int, start_row: int, end_row: int) -> bool:
        first_text = table.cell(start_row, col).text.strip()
        for row in range(start_row + 1, end_row + 1):
            if table.cell(row, col).text.strip() != first_text:
                return False
        return True

    def _merge_adjacent_equal_cells(self, table, col: int, start_row: int, end_row: int):
        if start_row >= end_row:
            return
        range_start = start_row
        range_value = table.cell(range_start, col).text.strip()
        for row in range(start_row + 1, end_row + 1):
            current_value = table.cell(row, col).text.strip()
            if current_value == range_value:
                continue
            else:
                if row - 1 > range_start:
                    for r in range(range_start + 1, row):
                        table.cell(r, col).text = ""
                    self._merge_cells_vertical(table, col, range_start, row - 1)
                range_start = row
                range_value = current_value
        if end_row > range_start:
            for r in range(range_start + 1, end_row + 1):
                table.cell(r, col).text = ""
            self._merge_cells_vertical(table, col, range_start, end_row)

    def _merge_column_2_by_level(self, table, start_row, actual_start, actual_end, rows_data):
        """Объединяет колонку 2 (№ инструкции) по (maintenance_type, subgroup)."""
        if actual_end <= actual_start:
            return

        keys = []
        for row_idx in range(actual_start, actual_end + 1):
            row_data_index = row_idx - start_row
            if row_data_index < len(rows_data):
                row_data = rows_data[row_data_index]
                instr = row_data[2] if len(row_data) > 2 and row_data[0] == 'object' else None
                if instr:
                    maint = (instr.maintenance_type or "").lower()
                    subgroup = (instr.subgroup or "").strip() or "_default_"
                    keys.append((maint, subgroup))
                else:
                    keys.append(None)
            else:
                keys.append(None)

        unique_keys = set([k for k in keys if k is not None])
        if len(unique_keys) <= 1:
            self._merge_adjacent_equal_cells(table, 2, actual_start, actual_end)
        else:
            current_key = None
            sub_start = actual_start
            for row_idx in range(actual_start, actual_end + 1):
                row_data_index = row_idx - start_row
                key = None
                if row_data_index < len(rows_data):
                    row_data = rows_data[row_data_index]
                    instr = row_data[2] if len(row_data) > 2 and row_data[0] == 'object' else None
                    if instr:
                        maint = (instr.maintenance_type or "").lower()
                        subgroup = (instr.subgroup or "").strip() or "_default_"
                        key = (maint, subgroup)

                if key != current_key:
                    if current_key is not None:
                        if sub_start <= row_idx - 1:
                            self._merge_adjacent_equal_cells(table, 2, sub_start, row_idx - 1)
                    current_key = key
                    sub_start = row_idx
            if current_key is not None:
                if sub_start <= actual_end:
                    self._merge_adjacent_equal_cells(table, 2, sub_start, actual_end)

    def _get_category_priority(self, session) -> dict:
        categories = session.query(DBCategory).order_by(DBCategory.sort_order).all()
        return {cat.name: cat.sort_order for cat in categories}

    def _get_room_category_id(self, session, room_category_name: str) -> int:
        if not room_category_name:
            return None
        rc = session.query(RoomCategory).filter(RoomCategory.name == room_category_name).first()
        return rc.id if rc else None

    def _get_instruction_priority(self, instr, target_enterprise: str = None,
                                  target_room_name: str = None, target_room_category_id: int = None) -> int:
        ent_match = (instr.enterprise or "").lower().strip() == (target_enterprise or "").lower().strip()
        room_match = (instr.room_name or "").lower().strip() == (target_room_name or "").lower().strip()
        cat_match = (instr.room_category_id == target_room_category_id)

        if ent_match and room_match and cat_match:
            return 1
        elif ent_match and room_match:
            return 2
        elif ent_match and cat_match and (instr.room_name or "").strip() == "":
            return 3
        elif ent_match and (instr.room_name or "").strip() == "" and instr.room_category_id is None:
            return 4
        elif not ent_match and room_match and cat_match:
            return 5
        elif not ent_match and room_match:
            return 6
        elif not ent_match and cat_match and (instr.room_name or "").strip() == "":
            return 7
        elif not ent_match and (instr.room_name or "").strip() == "" and instr.room_category_id is None:
            return 8
        else:
            return 99

    def _filter_instructions(self, instructions: list, target_enterprise: str = None,
                             target_room_name: str = None, target_room_category_id: int = None) -> list:
        if not instructions:
            return []

        target_ent = (target_enterprise or "").strip()
        target_room = (target_room_name or "").strip()

        if not target_ent and not target_room:
            return [i for i in instructions if i.room_category_id == target_room_category_id or i.room_category_id is None]

        filtered = [i for i in instructions if i.room_category_id == target_room_category_id or i.room_category_id is None]
        if not filtered:
            return []

        both = [i for i in filtered if (i.enterprise or "").strip() == target_ent and (i.room_name or "").strip() == target_room]
        if both:
            return both

        ent_only = [i for i in filtered if (i.enterprise or "").strip() == target_ent and (i.room_name or "").strip() == ""]
        if ent_only:
            return ent_only

        general = [i for i in filtered if (i.enterprise or "").strip() == "" and (i.room_name or "").strip() == ""]
        return general

    def _filter_by_methods(self, instructions: list, target_enterprise: str = None,
                           target_room_name: str = None, target_room_category_id: int = None,
                           grouping: str = 'method') -> list:
        if not instructions:
            return []

        groups = defaultdict(list)
        for i in instructions:
            if grouping == 'method_subgroup':
                key = (i.cleaning_method or "", (i.subgroup or "").strip())
            elif grouping == 'method_surface':
                key = (i.cleaning_method or "", (i.surface_type or "").lower())
            else:
                key = (i.cleaning_method or "",)
            groups[key].append(i)

        result = []
        for key, group in groups.items():
            filtered_group = self._filter_instructions(
                group, target_enterprise, target_room_name, target_room_category_id
            )
            result.extend(filtered_group)
        return result

    def _get_support_additions(self, all_instrs, target_enterprise, room_category_id):
        target_ent = (target_enterprise or "").strip()
        if not target_ent or room_category_id is None:
            return []
        return [
            instr for instr in all_instrs
            if (instr.enterprise or "").strip().lower() == target_ent.lower()
            and instr.room_category_id == room_category_id
            and (instr.maintenance_type or "").lower() == "поддерживающая"
        ]

    def _append_support_to_list(self, instructions, support_additions):
        if not support_additions:
            return instructions
        if instructions is None:
            instructions = []
        existing_ids = {i.id for i in instructions}
        appended = []
        for instr in support_additions:
            if instr.id not in existing_ids:
                appended.append(instr)
        if not appended:
            return instructions
        return list(instructions) + appended

    def _append_support_to_multi(self, instructions, support_additions):
        if not support_additions:
            return instructions
        if instructions is None:
            instructions = []
        existing_ids = {i.id for i in instructions}
        appended = []
        for instr in support_additions:
            if instr.id not in existing_ids:
                appended.append(instr)
        if not appended:
            return instructions
        return list(instructions) + appended

    def _append_support_to_split(self, split_instructions, support_additions):
        if not support_additions:
            return split_instructions
        if split_instructions is None:
            split_instructions = {}
        for instr in support_additions:
            sf = (instr.surface_type or "").lower()
            if not sf:
                continue
            if sf not in split_instructions:
                split_instructions[sf] = []
            existing_ids = {i.id for i in split_instructions[sf]}
            if instr.id not in existing_ids:
                split_instructions[sf].append(instr)
        for sf in split_instructions:
            split_instructions[sf].sort(
                key=lambda x: (
                    self.LEVEL_ORDER.get((x.maintenance_type or "").lower(), 99),
                    self.cleaning_method_order.get((x.cleaning_method or "").lower().strip(), 99)
                )
            )
        return split_instructions

    def _get_instruction_signature(self, instructions: list) -> tuple:
        if not instructions:
            return (("empty",),)
        signatures = []
        sorted_instructions = sorted(
            instructions,
            key=lambda x: self.cleaning_method_order.get(
                (x.cleaning_method or "").lower().strip(), 99
            )
        )
        for instr in sorted_instructions:
            sig = (
                instr.maintenance_type or "",
                instr.cleaning_method or "",
                instr.instruction_number or "",
                instr.product_name or "",
                instr.cleaning_technique or "",
                instr.concentration or "",
                instr.temperature or "",
                instr.exposure_time or "",
                instr.inventory or "",
                instr.frequency or "",
                instr.executor or "",
                instr.control_method or "",
                instr.surface_type or "",
                instr.enterprise or "",
                instr.subgroup or "",
                instr.room_name or "",
                instr.room_category_id or 0,
            )
            signatures.append(sig)
        return tuple(signatures)

    def _group_by_full_instructions(self, object_instructions: list) -> list:
        groups = {}
        for obj_name, instructions in object_instructions:
            signature = self._get_instruction_signature(instructions)
            if signature not in groups:
                groups[signature] = {"names": [], "instructions": instructions}
            groups[signature]["names"].append(obj_name)
        result = []
        for signature, data in groups.items():
            merged_name = ", ".join(sorted(data["names"]))
            result.append((merged_name, data["instructions"]))
        return result

    def _select_instructions_for_room(self, all_instructions: list, room_category_id: int,
                                      enterprise: str = None, room_name: str = None,
                                      product_name: str = None, application_method: str = None) -> list:
        if not all_instructions:
            return []

        filtered = self._filter_instructions(all_instructions, enterprise, room_name, room_category_id)
        if not filtered:
            return []

        if product_name:
            normalized_product = self._normalize_product_name(product_name)
            if normalized_product:
                filtered = [
                    instr for instr in filtered
                    if instr.cleaning_method != "дезинфекция"
                    or self._normalize_product_name(instr.product_name or "") == normalized_product
                ]

        if application_method:
            filtered_before = list(filtered)
            filtered = [
                instr for instr in filtered
                if not instr.application_method or instr.application_method == application_method
            ]
            if not filtered:
                filtered = filtered_before

        by_method = defaultdict(list)
        for instr in filtered:
            method = instr.cleaning_method or ""
            by_method[method].append(instr)

        selected = []
        for method, instrs in by_method.items():
            sorted_instrs = sorted(
                instrs,
                key=lambda i: self._get_instruction_priority(i, enterprise, room_name, room_category_id)
            )
            best = None
            for maint_level in ["основная", "поддерживающая", "генеральная"]:
                for instr in sorted_instrs:
                    if (instr.maintenance_type or "").lower() == maint_level:
                        best = instr
                        break
                if best:
                    break
            if best:
                selected.append(best)
        selected.sort(
            key=lambda x: self.cleaning_method_order.get(
                (x.cleaning_method or "").lower().strip(), 99
            )
        )
        return selected

    def _select_all_instructions_for_room(self, all_instructions: list, room_category_id: int,
                                          enterprise: str = None, room_name: str = None,
                                          application_method: str = None) -> list:
        if not all_instructions:
            return []

        filtered = self._filter_instructions(all_instructions, enterprise, room_name, room_category_id)
        if not filtered:
            return []

        if application_method:
            filtered_before = list(filtered)
            filtered = [
                instr for instr in filtered
                if not instr.application_method or instr.application_method == application_method
            ]
            if not filtered:
                filtered = filtered_before

        sorted_instrs = sorted(
            filtered,
            key=lambda i: (
                self._get_instruction_priority(i, enterprise, room_name, room_category_id),
                self.LEVEL_ORDER.get((i.maintenance_type or "").lower(), 99),
                self.cleaning_method_order.get((i.cleaning_method or "").lower().strip(), 99)
            )
        )

        level_groups = defaultdict(list)
        for instr in sorted_instrs:
            level = (instr.maintenance_type or "").lower()
            level_groups[level].append(instr)

        result = []
        for level in ["основная", "поддерживающая", "генеральная"]:
            instrs = level_groups.get(level, [])
            if not instrs:
                continue
            method_groups = defaultdict(list)
            for instr in instrs:
                method = instr.cleaning_method or ""
                method_groups[method].append(instr)
            ordered_methods = sorted(method_groups.keys(),
                                     key=lambda m: self.cleaning_method_order.get(m.lower().strip(), 99))
            while True:
                added = False
                for method in ordered_methods:
                    if method_groups[method]:
                        result.append(method_groups[method].pop(0))
                        added = True
                if not added:
                    break
        return result

    def _select_split_instructions(self, all_instructions: list, room_category_id: int,
                                   enterprise: str = None, room_name: str = None,
                                   cleaning_application_method: str = None,
                                   disinfection_product_name: str = None,
                                   disinfection_application_method: str = None) -> dict:
        result = {}
        surface_types = ["внешняя", "внутренняя", "очистка от мин. отложений"]

        for sf in surface_types:
            sf_instrs = [i for i in all_instructions if (i.surface_type or "").lower() == sf]
            if not sf_instrs:
                continue

            disinfection_instrs = [i for i in sf_instrs if i.cleaning_method == "дезинфекция"]
            other_instrs = [i for i in sf_instrs if i.cleaning_method != "дезинфекция"]

            selected_other = []
            if other_instrs:
                selected_other = self._select_instructions_for_room(
                    other_instrs, room_category_id, enterprise, room_name,
                    product_name=None,
                    application_method=cleaning_application_method
                )

            selected_disinfection = None
            if disinfection_instrs:
                if disinfection_product_name:
                    normalized_product = self._normalize_product_name(disinfection_product_name)
                    candidates = [
                        instr for instr in disinfection_instrs
                        if not normalized_product
                        or self._normalize_product_name(instr.product_name or "") == normalized_product
                    ]
                    if candidates:
                        candidates.sort(
                            key=lambda i: self._get_instruction_priority(i, enterprise, room_name, room_category_id)
                        )
                        selected_disinfection = candidates[0]
                    else:
                        if disinfection_instrs:
                            disinfection_instrs_sorted = sorted(
                                disinfection_instrs,
                                key=lambda i: self._get_instruction_priority(i, enterprise, room_name, room_category_id)
                            )
                            selected_disinfection = disinfection_instrs_sorted[0]
                else:
                    disinfection_selected = self._select_instructions_for_room(
                        disinfection_instrs, room_category_id, enterprise, room_name,
                        product_name=None,
                        application_method=disinfection_application_method
                    )
                    if disinfection_selected:
                        selected_disinfection = disinfection_selected[0]

            selected = list(selected_other)
            if selected_disinfection:
                selected.append(selected_disinfection)

            if selected:
                selected.sort(
                    key=lambda x: self.cleaning_method_order.get(
                        (x.cleaning_method or "").lower().strip(), 99
                    )
                )
                result[sf] = selected

        return result

    def _clone_row_formatting(self, table, source_row_idx: int, target_row_idx: int):
        source_row = table.rows[source_row_idx]
        target_row = table.rows[target_row_idx]
        for col_idx, source_cell in enumerate(source_row.cells):
            if col_idx >= len(target_row.cells):
                break
            target_cell = target_row.cells[col_idx]
            source_tc_pr = source_cell._tc.find(qn('w:tcPr'))
            if source_tc_pr is not None:
                target_tc_pr = target_cell._tc.find(qn('w:tcPr'))
                if target_tc_pr is not None:
                    target_cell._tc.remove(target_tc_pr)
                target_cell._tc.insert(0, copy.deepcopy(source_tc_pr))
            for p_idx, source_para in enumerate(source_cell.paragraphs):
                if p_idx >= len(target_cell.paragraphs):
                    target_para = target_cell.add_paragraph()
                else:
                    target_para = target_cell.paragraphs[p_idx]
                source_p_pr = source_para._p.find(qn('w:pPr'))
                if source_p_pr is not None:
                    target_p_pr = target_para._p.find(qn('w:pPr'))
                    if target_p_pr is not None:
                        target_para._p.remove(target_p_pr)
                    target_para._p.insert(0, copy.deepcopy(source_p_pr))
                for run in target_para.runs:
                    target_para._p.remove(run._r)
                for source_run in source_para.runs:
                    new_run = copy.deepcopy(source_run._r)
                    r_pr = new_run.find(qn('w:rPr'))
                    if r_pr is None:
                        r_pr = OxmlElement('w:rPr')
                        new_run.insert(0, r_pr)
                    source_r_pr = source_run._r.find(qn('w:rPr'))
                    if source_r_pr is not None:
                        source_b = source_r_pr.find(qn('w:b'))
                        if source_b is not None:
                            existing_b = r_pr.find(qn('w:b'))
                            if existing_b is None:
                                r_pr.append(copy.deepcopy(source_b))
                    if col_idx == 3:
                        existing_b = r_pr.find(qn('w:b'))
                        if existing_b is None:
                            b = OxmlElement('w:b')
                            r_pr.append(b)
                    if col_idx == 8:
                        existing_b = r_pr.find(qn('w:b'))
                        if existing_b is None:
                            b = OxmlElement('w:b')
                            r_pr.append(b)
                    sz = r_pr.find(qn('w:sz'))
                    if sz is None:
                        sz = OxmlElement('w:sz')
                        r_pr.append(sz)
                    sz.set(qn('w:val'), '14')
                    sz_cs = r_pr.find(qn('w:szCs'))
                    if sz_cs is None:
                        sz_cs = OxmlElement('w:szCs')
                        r_pr.append(sz_cs)
                    sz_cs.set(qn('w:val'), '14')
                    r_fonts = r_pr.find(qn('w:rFonts'))
                    if r_fonts is None:
                        r_fonts = OxmlElement('w:rFonts')
                        r_pr.append(r_fonts)
                    r_fonts.set(qn('w:ascii'), 'Arial')
                    r_fonts.set(qn('w:hAnsi'), 'Arial')
                    r_fonts.set(qn('w:cs'), 'Arial')
                    for t in new_run.findall(qn('w:t')):
                        t.text = ''
                    target_para._p.append(new_run)

    def _set_cell_text(self, cell, text: str, bold: bool = False):
        text = self._clean_text(text)
        if not cell.paragraphs:
            cell.add_paragraph()
        para = cell.paragraphs[0]
        if not para.runs:
            run_elem = OxmlElement('w:r')
            r_pr = OxmlElement('w:rPr')
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '14')
            r_pr.append(sz)
            r_fonts = OxmlElement('w:rFonts')
            r_fonts.set(qn('w:ascii'), 'Arial')
            r_fonts.set(qn('w:hAnsi'), 'Arial')
            r_pr.append(r_fonts)
            if bold:
                b = OxmlElement('w:b')
                r_pr.append(b)
            run_elem.append(r_pr)
            t_elem = OxmlElement('w:t')
            t_elem.set(qn('xml:space'), 'preserve')
            run_elem.append(t_elem)
            para._p.append(run_elem)
        else:
            run = para.runs[0]
            r_pr = run._r.find(qn('w:rPr'))
            if r_pr is None:
                r_pr = OxmlElement('w:rPr')
                run._r.insert(0, r_pr)
            existing_b = r_pr.find(qn('w:b'))
            if existing_b is not None:
                r_pr.remove(existing_b)
            if bold:
                b = OxmlElement('w:b')
                r_pr.append(b)
        run = para.runs[0]
        for t in run._r.findall(qn('w:t')):
            t.text = text
            break

    def _set_product_cell(self, cell, text: str, bold: bool = False):
        text = self._clean_text(text)
        if not cell.paragraphs:
            cell.add_paragraph()
        para = cell.paragraphs[0]
        for r in para.runs:
            para._p.remove(r._r)
        parts = text.split('®')
        for i, part in enumerate(parts):
            if part:
                run = para.add_run(part)
                run.font.name = 'Arial'
                run.font.size = Pt(7)
                if bold:
                    run.bold = True
            if i < len(parts) - 1:
                run_reg = para.add_run('®')
                run_reg.font.name = 'Arial'
                run_reg.font.size = Pt(7)
                if bold:
                    run_reg.bold = True
                run_reg.font.superscript = True

    def _clear_cell_text(self, cell):
        for para in cell.paragraphs:
            for run in para.runs:
                for t in run._r.findall(qn('w:t')):
                    t.text = ''

    def _is_valid_product_name(self, name: str) -> bool:
        if not name:
            return False
        name_lower = name.lower().strip()
        if name_lower in ('выберите элемент.', 'не выбрано', 'выберите элемент', ''):
            return False
        patterns = [
            r'%', r'дозирующая система', r'ручной', r'система', r'протирание',
            r'щётка', r'ветошь', r'губка', r'погружение', r'замачивание',
            r'орошение', r'распыление', r'протирка', r'обработка', r'раствор'
        ]
        for pattern in patterns:
            if re.search(pattern, name_lower):
                return False
        if re.match(r'^[0-9.,%\-–\s]+$', name_lower):
            return False
        return True

    def generate(self, checklist_data: ChecklistData, output_path: str, mode: int = 1,
                 enterprise_products_path: str = None, progress_callback=None) -> str:
        try:
            if not self.template_path.exists():
                raise FileNotFoundError(f"Шаблон не найден: {self.template_path}")

            doc = Document(self.template_path)
            if len(doc.tables) < 2:
                raise ValueError("В шаблоне должно быть как минимум 2 таблицы")
            main_table = doc.tables[1]

            room_cell = main_table.cell(1, 0)
            if "Помещение:" in room_cell.text:
                room_cell.text = f"Помещение: {checklist_data.room_name or '___________'}"
                for para in room_cell.paragraphs:
                    for run in para.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9)
                        run.font.bold = True

            warning_row = main_table.rows[3]
            if checklist_data.room_category not in ["производственное", "складское"]:
                self._clear_cell_text(warning_row.cells[0])

            session = SessionLocal()
            self._load_config(session)

            category_priority = self._get_category_priority(session)
            product_colors = self._load_product_colors(session)

            target_enterprise = checklist_data.enterprise
            target_room_name = checklist_data.room_name

            room_category_id = None
            if hasattr(checklist_data, 'room_category') and checklist_data.room_category:
                room_category_id = self._get_room_category_id(session, checklist_data.room_category)

            checked_items = checklist_data.get_checked_items()
            all_names = [item.name for item in checked_items]

            db_objects = session.query(DBObject).filter(DBObject.normalized_name.in_(all_names)).all()
            objects_dict = {obj.normalized_name: obj for obj in db_objects}

            object_ids = [obj.id for obj in db_objects]
            db_instructions = session.query(Instruction).filter(Instruction.object_id.in_(object_ids)).all()

            instructions_dict = defaultdict(list)
            for instr in db_instructions:
                instructions_dict[instr.object_id].append(instr)

            category_object_instructions = defaultdict(list)
            all_object_instructions = []
            category_order = {}
            unmatched_objects = []
            unmatched_reasons = {}

            for item in checked_items:
                obj = objects_dict.get(item.name)
                cat_name = item.category.value
                priority = category_priority.get(cat_name, 999)
                category_order[cat_name] = priority

                display_name = obj.display_name if obj else item.name
                sort_priority = obj.sort_priority if obj else 999
                normalized_name = item.name

                print(f"\n🔍 Обработка объекта: {normalized_name}")
                print(f"   sort_priority = {sort_priority}")

                if obj:
                    print(f"   ID объекта в БД: {obj.id}")
                    all_instrs = instructions_dict.get(obj.id, [])
                    print(f"   Всего инструкций в БД: {len(all_instrs)}")
                else:
                    print(f"   ⚠️ Объект не найден в БД")
                    unmatched_objects.append(display_name)
                    unmatched_reasons[display_name] = "Объект не найден в БД"
                    continue

                if not (obj and obj.id in instructions_dict):
                    unmatched_objects.append(display_name)
                    unmatched_reasons[display_name] = "объект не найден в instructions_dict"
                    continue

                all_instrs = instructions_dict[obj.id]

                is_multi = normalized_name in self.multi_method_objects
                is_split = normalized_name in self.split_objects
                is_support = normalized_name in self.support_objects

                if is_support:
                    print(f"   -> объект в support_objects (поддерживающая — надстройка)")

                cleaning_application_method = self._get_checklist_method(
                    'мойка', normalized_name, checklist_data
                )
                if cleaning_application_method:
                    print(f"   💧 Способ разведения для мойки: '{cleaning_application_method}'")

                maintenance_instrs = [
                    i for i in all_instrs
                    if (i.maintenance_type or "").lower() == "поддерживающая"
                ]
                non_maintenance_instrs = [
                    i for i in all_instrs
                    if (i.maintenance_type or "").lower() != "поддерживающая"
                ]

                print(f"   Разделение: не-поддерживающих {len(non_maintenance_instrs)}, поддерживающих {len(maintenance_instrs)}")

                if is_multi:
                    grouping = 'method_subgroup'
                elif is_split:
                    grouping = 'method_surface'
                else:
                    grouping = 'method'

                filtered_instrs = self._filter_by_methods(
                    non_maintenance_instrs, target_enterprise, target_room_name,
                    room_category_id, grouping=grouping
                )
                print(f"   После фильтрации ({grouping}): {len(filtered_instrs)}")

                result_instructions = None
                result_type = 'normal'

                if is_multi:
                    print(f"   -> объект в multi_method_objects")
                    # Отсев инструкций с заполненным surface_type
                    filtered_instrs = [i for i in filtered_instrs if not (i.surface_type or "").strip()]
                    print(f"   После отсева по surface_type: {len(filtered_instrs)}")

                    disinfection_instrs = [i for i in filtered_instrs if i.cleaning_method == "дезинфекция"]
                    other_instrs = [i for i in filtered_instrs if i.cleaning_method != "дезинфекция"]

                    other_selected = []
                    if other_instrs:
                        if cleaning_application_method:
                            other_selected = [
                                i for i in other_instrs
                                if not i.application_method
                                or i.application_method == cleaning_application_method
                            ]
                            if not other_selected:
                                other_selected = list(other_instrs)
                        else:
                            other_selected = list(other_instrs)

                    disinfection_selected = []
                    if disinfection_instrs:
                        normalized_product = None
                        if checklist_data.disinfection_product:
                            normalized_product = self._normalize_product_name(checklist_data.disinfection_product)

                        subgroup_groups = defaultdict(list)
                        for instr in disinfection_instrs:
                            subgroup_groups[(instr.subgroup or "").strip()].append(instr)

                        for sg, group in subgroup_groups.items():
                            if normalized_product:
                                candidates = [
                                    instr for instr in group
                                    if self._normalize_product_name(instr.product_name or "") == normalized_product
                                ]
                                if candidates:
                                    disinfection_selected.extend(candidates)
                                    continue
                            disinfection_selected.extend(group)

                        disinfection_selected.sort(
                            key=lambda i: (
                                self._get_instruction_priority(i, target_enterprise, target_room_name, room_category_id),
                                self.LEVEL_ORDER.get((i.maintenance_type or "").lower(), 99),
                                (i.subgroup or "")
                            )
                        )

                    instructions = other_selected + disinfection_selected
                    instructions.sort(
                        key=lambda x: (
                            (x.subgroup or ""),
                            self.cleaning_method_order.get((x.cleaning_method or "").lower().strip(), 99)
                        )
                    )
                    seen = set()
                    unique_instrs = []
                    for instr in instructions:
                        key = (instr.cleaning_method, instr.product_name, instr.surface_type, instr.subgroup)
                        if key not in seen:
                            seen.add(key)
                            unique_instrs.append(instr)
                    result_instructions = unique_instrs
                    result_type = 'normal'

                elif is_split:
                    print(f"   -> объект в split_objects")
                    result_instructions = self._select_split_instructions(
                        filtered_instrs, room_category_id,
                        target_enterprise, target_room_name,
                        cleaning_application_method=cleaning_application_method,
                        disinfection_product_name=checklist_data.disinfection_product,
                        disinfection_application_method=checklist_data.disinfection_method_text
                    )
                    result_type = 'split'

                else:
                    print(f"   -> обычный объект")
                    disinfection_instrs = [i for i in filtered_instrs if i.cleaning_method == "дезинфекция"]
                    other_instrs = [i for i in filtered_instrs if i.cleaning_method != "дезинфекция"]

                    selected_disinfection = None
                    if checklist_data.disinfection_product:
                        normalized_product = self._normalize_product_name(checklist_data.disinfection_product)
                        candidates = [i for i in disinfection_instrs
                                      if self._normalize_product_name(i.product_name or "") == normalized_product]
                        if candidates:
                            candidates.sort(key=lambda i: self._get_instruction_priority(i, target_enterprise, target_room_name, room_category_id))
                            selected_disinfection = candidates[0]
                        else:
                            if disinfection_instrs:
                                disinfection_instrs.sort(key=lambda i: self._get_instruction_priority(i, target_enterprise, target_room_name, room_category_id))
                                selected_disinfection = disinfection_instrs[0]
                    else:
                        if disinfection_instrs:
                            disinfection_selected_list = self._select_instructions_for_room(
                                disinfection_instrs, room_category_id,
                                target_enterprise, target_room_name,
                                product_name=None, application_method=None
                            )
                            if disinfection_selected_list:
                                selected_disinfection = disinfection_selected_list[0]

                    other_selected = []
                    if other_instrs:
                        other_selected = self._select_instructions_for_room(
                            other_instrs, room_category_id,
                            target_enterprise, target_room_name,
                            application_method=cleaning_application_method
                        )

                    instructions = other_selected
                    if selected_disinfection:
                        instructions.append(selected_disinfection)

                    instructions.sort(key=lambda x: self.cleaning_method_order.get((x.cleaning_method or "").lower().strip(), 99))

                    seen = set()
                    unique_instrs = []
                    for instr in instructions:
                        key = (instr.cleaning_method, instr.product_name, instr.surface_type)
                        if key not in seen:
                            seen.add(key)
                            unique_instrs.append(instr)
                    result_instructions = unique_instrs
                    result_type = 'normal'

                if is_support:
                    support_additions = self._get_support_additions(
                        all_instrs, target_enterprise, room_category_id
                    )
                    if support_additions:
                        print(f"   -> support: добавлено {len(support_additions)} поддерживающих инструкций")
                        if result_type == 'split':
                            result_instructions = self._append_support_to_split(
                                result_instructions, support_additions
                            )
                        elif is_multi:
                            result_instructions = self._append_support_to_multi(
                                result_instructions, support_additions
                            )
                        else:
                            result_instructions = self._append_support_to_list(
                                result_instructions, support_additions
                            )
                    else:
                        print(f"   -> support: поддерживающих инструкций не найдено (надстройка не применяется)")

                if result_type == 'split':
                    non_empty = {sf: instrs for sf, instrs in (result_instructions or {}).items() if instrs}
                    if non_empty:
                        category_object_instructions[cat_name].append(
                            ('split', display_name, non_empty, normalized_name, item))
                        all_object_instructions.append(
                            (display_name, non_empty, sort_priority, normalized_name, item))
                    else:
                        print(f"   ❌ split-инструкции не выбраны")
                        unmatched_objects.append(display_name)
                        unmatched_reasons[display_name] = "split-инструкции не выбраны"
                else:
                    if result_instructions:
                        print(f"   итог: {[f'{i.cleaning_method}/{i.maintenance_type} (ID={i.id})' for i in result_instructions]}")
                        category_object_instructions[cat_name].append(
                            ('normal', display_name, result_instructions, normalized_name, item))
                        all_object_instructions.append(
                            (display_name, result_instructions, sort_priority, normalized_name, item))
                    else:
                        print(f"   ❌ инструкции не выбраны для {normalized_name}")
                        unmatched_objects.append(display_name)
                        unmatched_reasons[display_name] = "инструкции не выбраны (пустой набор)"

            session.close()

            if unmatched_objects:
                print(f"\n⚠️ Объекты без инструкций: {unmatched_objects}")

            rows_data = []
            merge_info_object = []
            merge_info_columns = []
            surface_merge_info = []

            if mode == 1:
                sorted_categories = sorted(category_object_instructions.keys(), key=lambda x: category_order.get(x, 999))
                for cat_name in sorted_categories:
                    items = category_object_instructions[cat_name]
                    if not items:
                        continue
                    items.sort(key=lambda x: x[1])
                    rows_data.append(('category', cat_name, None))
                    for typ, display_name, instructions, normalized_name, item in items:
                        if typ == 'split':
                            surfaces = ["внешняя", "внутренняя", "очистка от мин. отложений"]
                            first_surface = True
                            group_start_row = len(rows_data)
                            for sf in surfaces:
                                if sf not in instructions:
                                    continue
                                sf_instrs = instructions[sf]
                                sf_label = {"внешняя": "Внешние поверхности",
                                            "внутренняя": "Внутренние поверхности",
                                            "очистка от мин. отложений": "Очистка от минеральных отложений"}[sf]
                                if first_surface:
                                    rows_data.append(('section_header', sf_label, display_name))
                                    first_surface = False
                                else:
                                    rows_data.append(('section_header', sf_label, None))
                                surface_start = len(rows_data)
                                for instr in sf_instrs:
                                    rows_data.append(('object', "", instr, normalized_name, instr.subgroup, item, sort_priority))
                                surface_end = len(rows_data) - 1
                                if surface_end >= surface_start:
                                    surface_merge_info.append((surface_start, surface_end))
                            group_end_row = len(rows_data) - 1
                            merge_info_object.append((group_start_row, group_end_row))
                        else:
                            group_start_row = len(rows_data)
                            if instructions:
                                for i, instr in enumerate(instructions):
                                    cell_text = display_name if i == 0 else ""
                                    rows_data.append(('object', cell_text, instr, normalized_name, instr.subgroup, item, sort_priority))
                            else:
                                rows_data.append(('object', display_name, None, normalized_name, None, item, sort_priority))
                            group_end_row = len(rows_data) - 1
                            if group_end_row > group_start_row:
                                merge_info_object.append((group_start_row, group_end_row))
                                merge_info_columns.append((group_start_row, group_end_row))
            elif mode == 2:
                all_object_instructions.sort(key=lambda x: (x[2], x[0]))
                for display_name, instructions, sort_priority, normalized_name, item in all_object_instructions:
                    if isinstance(instructions, dict):
                        surfaces = ["внешняя", "внутренняя", "очистка от мин. отложений"]
                        first_surface = True
                        group_start_row = len(rows_data)
                        for sf in surfaces:
                            if sf not in instructions:
                                continue
                            sf_instrs = instructions[sf]
                            sf_label = {"внешняя": "Внешние поверхности",
                                        "внутренняя": "Внутренние поверхности",
                                        "очистка от мин. отложений": "Очистка от минеральных отложений"}[sf]
                            if first_surface:
                                rows_data.append(('section_header', sf_label, display_name))
                                first_surface = False
                            else:
                                rows_data.append(('section_header', sf_label, None))
                            surface_start = len(rows_data)
                            for instr in sf_instrs:
                                rows_data.append(('object', "", instr, normalized_name, instr.subgroup, item, sort_priority))
                            surface_end = len(rows_data) - 1
                            if surface_end >= surface_start:
                                surface_merge_info.append((surface_start, surface_end))
                        group_end_row = len(rows_data) - 1
                        merge_info_object.append((group_start_row, group_end_row))
                    else:
                        group_start_row = len(rows_data)
                        if instructions:
                            for i, instr in enumerate(instructions):
                                cell_text = display_name if i == 0 else ""
                                rows_data.append(('object', cell_text, instr, normalized_name, instr.subgroup, item, sort_priority))
                        else:
                            rows_data.append(('object', display_name, None, normalized_name, None, item, sort_priority))
                        group_end_row = len(rows_data) - 1
                        if group_end_row > group_start_row:
                            merge_info_object.append((group_start_row, group_end_row))
                            merge_info_columns.append((group_start_row, group_end_row))
            elif mode == 3:
                inserted_headers = set()
                priority_groups = defaultdict(list)
                for display_name, instructions, sort_priority, normalized_name, item in all_object_instructions:
                    priority_groups[sort_priority].append((display_name, instructions, normalized_name, item))

                for priority in sorted(priority_groups.keys()):
                    items = priority_groups[priority]
                    normal_items = [(dn, instr, nn, it) for dn, instr, nn, it in items if not isinstance(instr, dict)]
                    split_items = [(dn, instr, nn, it) for dn, instr, nn, it in items if isinstance(instr, dict)]

                    groups = {}
                    for dn, instr, nn, it in normal_items:
                        signature = self._get_instruction_signature(instr)
                        if signature not in groups:
                            groups[signature] = {"names": [], "instructions": instr, "normalized_name": nn, "item": it}
                        groups[signature]["names"].append(dn)

                    for sig, data in groups.items():
                        current_nn = data["normalized_name"]
                        item = data["item"]
                        for group_set, header_text in self.group_headers.items():
                            if current_nn in group_set and group_set not in inserted_headers:
                                rows_data.append(('group_header', header_text, None))
                                inserted_headers.add(group_set)
                                break
                        merged_name = ", ".join(sorted(data["names"]))
                        group_start_row = len(rows_data)
                        if data["instructions"]:
                            for i, instr in enumerate(data["instructions"]):
                                cell_text = merged_name if i == 0 else ""
                                rows_data.append(
                                    ('object', cell_text, instr, data["normalized_name"], instr.subgroup, item, priority))
                        else:
                            rows_data.append(('object', merged_name, None, data["normalized_name"], None, item, priority))
                        group_end_row = len(rows_data) - 1
                        if group_end_row > group_start_row:
                            merge_info_object.append((group_start_row, group_end_row))
                            merge_info_columns.append((group_start_row, group_end_row))

                    for dn, split_instr, nn, it in split_items:
                        for group_set, header_text in self.group_headers.items():
                            if nn in group_set and group_set not in inserted_headers:
                                rows_data.append(('group_header', header_text, None))
                                inserted_headers.add(group_set)
                                break
                        surfaces = ["внешняя", "внутренняя", "очистка от мин. отложений"]
                        first_surface = True
                        group_start_row = len(rows_data)
                        for sf in surfaces:
                            if sf not in split_instr:
                                continue
                            sf_instrs = split_instr[sf]
                            sf_label = {"внешняя": "Внешние поверхности",
                                        "внутренняя": "Внутренние поверхности",
                                        "очистка от мин. отложений": "Очистка от минеральных отложений"}[sf]
                            if first_surface:
                                rows_data.append(('section_header', sf_label, dn))
                                first_surface = False
                            else:
                                rows_data.append(('section_header', sf_label, None))
                            surface_start = len(rows_data)
                            for instr in sf_instrs:
                                rows_data.append(('object', "", instr, nn, instr.subgroup, it, priority))
                            surface_end = len(rows_data) - 1
                            if surface_end >= surface_start:
                                surface_merge_info.append((surface_start, surface_end))
                        group_end_row = len(rows_data) - 1
                        merge_info_object.append((group_start_row, group_end_row))

            if unmatched_objects:
                rows_data.append(('category', 'Объекты без инструкций (требуют настройки)', None))
                for name in sorted(unmatched_objects):
                    rows_data.append(('object', name, None, None, None, None, None))

            start_row = 8
            tbl = main_table._tbl
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            tr_elements = tbl.findall('.//w:tr', namespaces=ns)
            while len(tr_elements) > start_row + 1:
                tbl.remove(tr_elements[-1])
                tr_elements = tbl.findall('.//w:tr', namespaces=ns)

            for _ in range(len(rows_data)):
                main_table.add_row()
                self._clone_row_formatting(main_table, start_row, len(main_table.rows) - 1)

            row_to_remove = main_table.rows[start_row]
            tbl.remove(row_to_remove._tr)

            current_row = start_row
            for row_info in rows_data:
                row = main_table.rows[current_row]

                if row_info[0] == 'category':
                    self._merge_cells_horizontal(row, 0, 11)
                    self._set_cell_text(row.cells[0], row_info[1])
                    if row.cells[0].paragraphs and row.cells[0].paragraphs[0].runs:
                        row.cells[0].paragraphs[0].runs[0].font.bold = True

                elif row_info[0] == 'group_header':
                    self._merge_cells_horizontal(row, 0, 11)
                    self._set_cell_text(row.cells[0], row_info[1], bold=True)
                    self._set_cell_background(row.cells[0], "D2D2D2")
                    for para in row.cells[0].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(7)

                elif row_info[0] == 'section_header':
                    if len(row_info) > 2 and row_info[2] is not None:
                        self._set_cell_text(row.cells[0], row_info[2])
                    else:
                        self._set_cell_text(row.cells[0], "")
                    self._merge_cells_horizontal(row, 1, 11)
                    self._set_cell_text(row.cells[1], row_info[1], bold=True)
                    self._set_cell_background(row.cells[1], "D2D2D2")
                    for para in row.cells[1].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(7)

                elif row_info[0] == 'object':
                    obj_name, instr = row_info[1], row_info[2]
                    normalized_name = row_info[3] if len(row_info) > 3 else None
                    item = row_info[5] if len(row_info) > 5 else None
                    sort_priority = row_info[6] if len(row_info) > 6 else 999

                    self._set_cell_text(row.cells[0], obj_name)

                    if instr:
                        cleaning_method = self._clean_text(instr.cleaning_method or "")
                        cleaning_technique = self._clean_text(instr.cleaning_technique or "")

                        db_product = self._clean_text(instr.product_name or "")
                        db_concentration = self._clean_text(instr.concentration or "")
                        db_method = self._clean_text(instr.application_method or "")

                        is_support_instr = (
                            normalized_name in self.support_objects
                            and (instr.maintenance_type or "").lower() == "поддерживающая"
                        )

                        if is_support_instr:
                            final_product = db_product if db_product else ""
                            final_concentration = db_concentration if db_concentration else "___________________"
                            final_extra_method = db_method if db_method else "___________"
                            final_temperature = self._clean_text(instr.temperature or "") if instr.temperature else "___________"
                            final_exposure = self._clean_text(instr.exposure_time or "") if instr.exposure_time else "___________"
                            final_technique = cleaning_technique if cleaning_technique else "___________"
                            final_control = self._clean_text(instr.control_method or "") if instr.control_method else "___________"
                            final_instruction_number = self._clean_text(instr.instruction_number or "")
                            final_frequency = self._clean_text(instr.frequency or "")

                            self._set_cell_text(row.cells[1], cleaning_method)
                            self._set_cell_text(row.cells[2], final_instruction_number)
                            if final_product:
                                self._set_product_cell(row.cells[3], final_product, bold=True)
                            else:
                                self._set_product_cell(row.cells[3], "_____________________", bold=False)
                            self._set_cell_text(row.cells[4], final_technique)

                            para = row.cells[5].paragraphs[0]
                            for r in para.runs:
                                para._p.remove(r._r)
                            run_conc = para.add_run(final_concentration)
                            run_conc.font.name = 'Arial'
                            run_conc.font.size = Pt(7)
                            if final_extra_method:
                                run_br = para.add_run()
                                br = OxmlElement('w:br')
                                run_br._r.append(br)
                                run_method = para.add_run(final_extra_method)
                                run_method.font.name = 'Arial'
                                run_method.font.size = Pt(7)

                            self._set_cell_text(row.cells[6], final_temperature)
                            self._set_cell_text(row.cells[7], final_exposure)

                            inv_text = self._clean_text(instr.inventory or "")
                            is_valid_color = inv_text and inv_text.lower() != "промаркированный" and inv_text.lower() in self._inventory_colors
                            if is_valid_color:
                                inv_color = self._inventory_colors[inv_text.lower()]
                                self._set_cell_text(row.cells[8], inv_text, bold=True)
                                self._set_cell_background(row.cells[8], inv_color)
                            elif inv_text:
                                self._set_cell_text(row.cells[8], inv_text, bold=True)
                            else:
                                self._set_cell_text(row.cells[8], "промаркированный", bold=True)

                            self._set_cell_text(row.cells[9], final_frequency)

                            executor_value = None
                            if instr.executor:
                                executor_value = instr.executor
                            else:
                                if item and isinstance(item, ChecklistItem):
                                    equipment_categories = {
                                        Category.THERMAL_EQUIPMENT, Category.TECH_EQUIPMENT,
                                        Category.REFRIGERATION_EQUIPMENT, Category.DISHWASHING_EQUIPMENT,
                                        Category.PACKAGING_EQUIPMENT, Category.DOSING_EQUIPMENT
                                    }
                                    is_equipment = (item.category in equipment_categories)
                                    if not is_equipment:
                                        if sort_priority == -1:
                                            executor_value = checklist_data.executor_high
                                        elif 1 <= sort_priority <= 50:
                                            executor_value = checklist_data.executor_low
                                    else:
                                        executor_value = checklist_data.executor_equipment
                            if not executor_value:
                                executor_value = ""
                            self._set_cell_text(row.cells[10], self._clean_text(executor_value))

                            self._set_cell_text(row.cells[11], final_control)

                            if final_product:
                                color = product_colors.get(final_product)
                                if color:
                                    self._set_cell_background(row.cells[3], color)

                            current_row += 1
                            continue

                        has_db_product = db_product and self._is_valid_product_name(db_product)
                        is_disinfection = (cleaning_method == "дезинфекция")

                        if is_disinfection:
                            final_product = db_product if db_product else ""
                            if db_concentration:
                                final_concentration = db_concentration
                            else:
                                checklist_concentration = checklist_data.disinfection_concentration
                                final_concentration = self._clean_text(checklist_concentration) if checklist_concentration else ""
                            if db_method:
                                final_extra_method = db_method
                            else:
                                checklist_method_text = checklist_data.disinfection_method_text
                                final_extra_method = self._clean_text(checklist_method_text) if checklist_method_text else ""
                            final_temperature = self._clean_text(instr.temperature or "") if instr.temperature else "___________"
                            final_exposure = self._clean_text(instr.exposure_time or "") if instr.exposure_time else "___________"

                            source_application_method = db_method if db_method else checklist_data.disinfection_method_text
                            if source_application_method:
                                method_text = source_application_method.lower()
                                target_method_name = None
                                if "promax" in method_text:
                                    target_method_name = "протирание"
                                elif "protwin" in method_text:
                                    target_method_name = "орошение"
                                elif "пенная станция" in method_text:
                                    target_method_name = "запенивание"
                                if target_method_name:
                                    found = self._cleaning_techniques.get(target_method_name.lower())
                                    if found:
                                        cleaning_technique = found

                            self._set_cell_text(row.cells[1], cleaning_method)
                            self._set_cell_text(row.cells[2], self._clean_text(instr.instruction_number or ""))
                            if final_product:
                                self._set_product_cell(row.cells[3], final_product, bold=True)
                            else:
                                self._set_product_cell(row.cells[3], "_____________________", bold=False)
                            self._set_cell_text(row.cells[4], cleaning_technique)

                            para = row.cells[5].paragraphs[0]
                            for r in para.runs:
                                para._p.remove(r._r)
                            conc_text = final_concentration if final_concentration else "___________________"
                            run_conc = para.add_run(conc_text)
                            run_conc.font.name = 'Arial'
                            run_conc.font.size = Pt(7)
                            if final_extra_method:
                                run_br = para.add_run()
                                br = OxmlElement('w:br')
                                run_br._r.append(br)
                                run_method = para.add_run(final_extra_method)
                                run_method.font.name = 'Arial'
                                run_method.font.size = Pt(7)

                            self._set_cell_text(row.cells[6], final_temperature)
                            self._set_cell_text(row.cells[7], final_exposure)

                        else:
                            if has_db_product:
                                final_product = db_product
                                if db_concentration:
                                    final_concentration = db_concentration
                                else:
                                    final_concentration = self._get_checklist_concentration(cleaning_method, normalized_name, checklist_data)
                                if db_method:
                                    final_extra_method = db_method
                                else:
                                    final_extra_method = self._get_checklist_method(cleaning_method, normalized_name, checklist_data)
                            else:
                                final_product = self._get_checklist_product(cleaning_method, normalized_name, checklist_data)
                                final_concentration = self._get_checklist_concentration(cleaning_method, normalized_name, checklist_data)
                                final_extra_method = self._get_checklist_method(cleaning_method, normalized_name, checklist_data)

                            if not final_product:
                                final_product = None

                            self._set_cell_text(row.cells[1], cleaning_method)
                            self._set_cell_text(row.cells[2], self._clean_text(instr.instruction_number or ""))
                            if final_product:
                                self._set_product_cell(row.cells[3], final_product, bold=True)
                            else:
                                self._set_product_cell(row.cells[3], "_____________________", bold=False)

                            self._set_cell_text(row.cells[4], cleaning_technique)

                            para = row.cells[5].paragraphs[0]
                            for r in para.runs:
                                para._p.remove(r._r)
                            conc_text = final_concentration if final_concentration else "___________________"
                            run_conc = para.add_run(conc_text)
                            run_conc.font.name = 'Arial'
                            run_conc.font.size = Pt(7)
                            if final_extra_method:
                                run_br = para.add_run()
                                br = OxmlElement('w:br')
                                run_br._r.append(br)
                                run_method = para.add_run(final_extra_method)
                                run_method.font.name = 'Arial'
                                run_method.font.size = Pt(7)

                            self._set_cell_text(row.cells[6], self._clean_text(instr.temperature or "") if instr.temperature else "___________")
                            self._set_cell_text(row.cells[7], self._clean_text(instr.exposure_time or "") if instr.exposure_time else "___________")

                        inv_text = self._clean_text(instr.inventory or "")
                        is_valid_color = inv_text and inv_text.lower() != "промаркированный" and inv_text.lower() in self._inventory_colors
                        if is_valid_color:
                            inv_color = self._inventory_colors[inv_text.lower()]
                            self._set_cell_text(row.cells[8], inv_text, bold=True)
                            self._set_cell_background(row.cells[8], inv_color)
                        else:
                            if checklist_data.inventory_color:
                                inv_color = self._inventory_colors.get(checklist_data.inventory_color)
                                if inv_color:
                                    self._set_cell_text(row.cells[8], checklist_data.inventory_color, bold=True)
                                    self._set_cell_background(row.cells[8], inv_color)
                                else:
                                    self._set_cell_text(row.cells[8], "промаркированный", bold=True)
                            else:
                                self._set_cell_text(row.cells[8], "промаркированный", bold=True)

                        self._set_cell_text(row.cells[9], self._clean_text(instr.frequency or ""))

                        executor_value = None
                        if instr.executor:
                            executor_value = instr.executor
                        else:
                            if item and isinstance(item, ChecklistItem):
                                equipment_categories = {
                                    Category.THERMAL_EQUIPMENT, Category.TECH_EQUIPMENT,
                                    Category.REFRIGERATION_EQUIPMENT, Category.DISHWASHING_EQUIPMENT,
                                    Category.PACKAGING_EQUIPMENT, Category.DOSING_EQUIPMENT
                                }
                                is_equipment = (item.category in equipment_categories)
                                if not is_equipment:
                                    if sort_priority == -1:
                                        executor_value = checklist_data.executor_high
                                    elif 1 <= sort_priority <= 50:
                                        executor_value = checklist_data.executor_low
                                else:
                                    executor_value = checklist_data.executor_equipment
                        if not executor_value:
                            executor_value = ""
                        self._set_cell_text(row.cells[10], self._clean_text(executor_value))

                        self._set_cell_text(row.cells[11], self._clean_text(instr.control_method or ""))

                        if final_product:
                            color = product_colors.get(final_product)
                            if color:
                                self._set_cell_background(row.cells[3], color)

                current_row += 1

            for group_start, group_end in merge_info_object:
                actual_start = start_row + group_start
                actual_end = start_row + group_end
                if actual_end > actual_start:
                    for row in range(actual_start + 1, actual_end + 1):
                        main_table.cell(row, 0).text = ""
                    self._merge_cells_vertical(main_table, 0, actual_start, actual_end)

            for group_start, group_end in merge_info_columns:
                actual_start = start_row + group_start
                actual_end = start_row + group_end
                if actual_end > actual_start:
                    for col in [8, 9, 10, 11]:
                        self._merge_adjacent_equal_cells(main_table, col, actual_start, actual_end)

            for group_start, group_end in merge_info_columns:
                actual_start = start_row + group_start
                actual_end = start_row + group_end
                self._merge_column_2_by_level(main_table, start_row, actual_start, actual_end, rows_data)

            # ===== ИСПРАВЛЕНО: для split-объектов колонка 2 тоже объединяется
            # по (maintenance_type, subgroup), а не по равенству текста =====
            for group_start, group_end in surface_merge_info:
                actual_start = start_row + group_start
                actual_end = start_row + group_end
                if actual_end > actual_start:
                    self._merge_column_2_by_level(main_table, start_row, actual_start, actual_end, rows_data)
                    for col in [8, 9, 10, 11]:
                        self._merge_adjacent_equal_cells(main_table, col, actual_start, actual_end)

            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            if output_file.exists():
                output_file.unlink()
            doc.save(str(output_file))
            return str(output_file)

        except Exception as e:
            print("=" * 80)
            print("❌ ОШИБКА В ГЕНЕРАТОРЕ:")
            traceback.print_exc()
            print("=" * 80)
            raise

    def _get_checklist_product(self, cleaning_method, normalized_name, checklist_data):
        if cleaning_method == "мойка":
            if normalized_name in self.floor_objects:
                return checklist_data.floor_cleaning_product
            elif normalized_name in self.tech_objects:
                return checklist_data.tech_cleaning_product
            elif normalized_name in self.thermal_objects:
                return checklist_data.thermal_cleaning_product
            elif normalized_name in self.glass_objects:
                return checklist_data.glass_cleaning_product
            else:
                return checklist_data.cleaning_product
        return None

    def _get_checklist_concentration(self, cleaning_method, normalized_name, checklist_data):
        if cleaning_method == "мойка":
            if normalized_name in self.floor_objects:
                return checklist_data.floor_cleaning_concentration
            elif normalized_name in self.tech_objects:
                return checklist_data.tech_cleaning_concentration
            elif normalized_name in self.thermal_objects:
                return checklist_data.thermal_cleaning_concentration
            elif normalized_name in self.glass_objects:
                return checklist_data.glass_cleaning_concentration
            else:
                return checklist_data.cleaning_concentration
        return None

    def _get_checklist_method(self, cleaning_method, normalized_name, checklist_data):
        if cleaning_method == "мойка":
            if normalized_name in self.floor_objects:
                return checklist_data.floor_cleaning_method_text
            elif normalized_name in self.tech_objects:
                return checklist_data.tech_cleaning_method_text
            elif normalized_name in self.thermal_objects:
                return checklist_data.thermal_cleaning_method_text
            elif normalized_name in self.glass_objects:
                return checklist_data.glass_cleaning_method_text
            else:
                return checklist_data.cleaning_method_text
        return None