from docx import Document
from pathlib import Path


def diagnose_template(template_path="tech_card_templates/шаблон.docx"):
    doc = Document(template_path)

    if not doc.tables:
        print("❌ В документе нет таблиц!")
        return

    table = doc.tables[0]
    print(f"✅ Таблица найдена")
    print(f"   Количество строк: {len(table.rows)}")

    # Проверяем первые 10 строк
    for row_idx, row in enumerate(table.rows[:10]):
        col_count = len(row.cells)
        print(f"   Строка {row_idx}: {col_count} колонок")
        if col_count > 0:
            first_cell_text = row.cells[0].text[:50] if row.cells[0].text else "(пусто)"
            print(f"      Первая ячейка: {first_cell_text}")

    # Находим строку с заголовками (где есть "Объект обработки")
    header_row_idx = None
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if "Объект обработки" in cell.text:
                header_row_idx = row_idx
                break
        if header_row_idx is not None:
            break

    if header_row_idx is not None:
        print(f"\n📌 Строка с заголовками: индекс {header_row_idx}")
        header_row = table.rows[header_row_idx]
        print(f"   Количество колонок: {len(header_row.cells)}")
        for i, cell in enumerate(header_row.cells):
            print(f"      Колонка {i}: {cell.text[:40]}")
    else:
        print("\n❌ Строка с заголовками 'Объект обработки' не найдена")


if __name__ == "__main__":
    diagnose_template()