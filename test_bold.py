# test_bold.py
import sys
from pathlib import Path

sys.path.insert(0, '.')

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import copy
from docx.oxml import OxmlElement

# Тест: создаём документ из шаблона, добавляем строку, клонируем, заполняем
doc = Document("tech_card_templates/шаблон.docx")
table = doc.tables[1]

# Клонируем строку 7 в новую
source_row = table.rows[7]
new_row = table.add_row()

for col_idx, source_cell in enumerate(source_row.cells):
    target_cell = new_row.cells[col_idx]

    # Копируем tcPr
    source_tc_pr = source_cell._tc.find(qn('w:tcPr'))
    if source_tc_pr is not None:
        target_tc_pr = target_cell._tc.find(qn('w:tcPr'))
        if target_tc_pr is not None:
            target_cell._tc.remove(target_tc_pr)
        target_cell._tc.insert(0, copy.deepcopy(source_tc_pr))

    # Копируем параграфы
    for p_idx, source_para in enumerate(source_cell.paragraphs):
        if p_idx >= len(target_cell.paragraphs):
            target_para = target_cell.add_paragraph()
        else:
            target_para = target_cell.paragraphs[p_idx]

        source_p_pr = source_para._p.find(qn('w:pPr'))
        if source_p_pr is not None:
            target_p_pr = target_para._p.find(qn('w:pPr'))
            if target_p_pr is not None:
                target_para._p.remove(target_p_pr)
            target_para._p.insert(0, copy.deepcopy(source_p_pr))

        for run in target_para.runs:
            target_para._p.remove(run._r)

        for source_run in source_para.runs:
            new_run = copy.deepcopy(source_run._r)
            r_pr = new_run.find(qn('w:rPr'))
            if r_pr is None:
                r_pr = OxmlElement('w:rPr')
                new_run.insert(0, r_pr)

            # Принудительный bold для колонки 4
            if col_idx == 3:
                existing_b = r_pr.find(qn('w:b'))
                if existing_b is not None:
                    r_pr.remove(existing_b)
                b = OxmlElement('w:b')
                r_pr.append(b)

            # Шрифт
            sz = r_pr.find(qn('w:sz'))
            if sz is None:
                sz = OxmlElement('w:sz')
                r_pr.append(sz)
            sz.set(qn('w:val'), '14')

            r_fonts = r_pr.find(qn('w:rFonts'))
            if r_fonts is None:
                r_fonts = OxmlElement('w:rFonts')
                r_pr.append(r_fonts)
            r_fonts.set(qn('w:ascii'), 'Arial')
            r_fonts.set(qn('w:hAnsi'), 'Arial')

            for t in new_run.findall(qn('w:t')):
                t.text = 'TEST'
                break

            target_para._p.append(new_run)

# Проверяем XML
cell = new_row.cells[3]
for run in cell.paragraphs[0].runs:
    r_pr = run._r.find(qn('w:rPr'))
    bold = r_pr.find(qn('w:b')) if r_pr is not None else None
    print(f"Run text='{run.text}', bold={'✅' if bold is not None else '❌'}")

# Сохраняем для проверки
doc.save("test_bold_output.docx")
print("\n✅ Сохранено: test_bold_output.docx")