#!/usr/bin/env python3
"""
prepare_import_json.py

Скрипт для подготовки JSON-файла импорта в базу данных.
Читает:
- текстовый файл с объектами (полный перечень всех поверхностей из чек-листа.txt)
- Excel-файл с инструкциями (Объекты для базы данных.xlsx)

Создает:
- import_data.json (для импорта в БД)
"""

import json
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass, field, asdict
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    print("Установите openpyxl: pip install openpyxl")
    exit(1)

# ============================================================
# СПИСОК ОБЪЕКТОВ С МОДИФИКАТОРАМИ
# ============================================================

MODIFIER_OBJECTS = {
    # Базовый объект: (base_name, {модификатор: полное_имя})
    "потолок (высота _____м)": {
        "base_name": "потолок",
        "modifiers": {
            "П": "потолок (высота _____м) П",
            "ОК": "потолок (высота _____м) ОК"
        }
    },
    "вытяжные зонты": {
        "base_name": "вытяжные зонты",
        "modifiers": {
            "Н": "вытяжные зонты Н",
            "А": "вытяжные зонты А"
        }
    },
    "формы для выпечки": {
        "base_name": "формы для выпечки",
        "modifiers": {
            "С": "формы для выпечки С",
            "Н": "формы для выпечки Н",
            "А": "формы для выпечки А"
        }
    },
    "листы для выпечки": {
        "base_name": "листы для выпечки",
        "modifiers": {
            "Н": "листы для выпечки Н",
            "А": "листы для выпечки А"
        }
    },
    "съёмные детали оборудования": {
        "base_name": "съёмные детали оборудования",
        "modifiers": {
            "Н": "съёмные детали оборудования Н",
            "А": "съёмные детали оборудования А"
        }
    },
    "ПММ": {
        "base_name": "ПММ",
        "modifiers": {
            "купольная": "ПММ купольная",
            "туннельная": "ПММ туннельная"
        }
    },
    "камеры": {
        "base_name": "камеры",
        "modifiers": {
            "холд.": "камеры холд.",
            "мороз.": "камеры мороз.",
            "шок. замор.": "камеры шок. замор."
        }
    },
    "плиты": {
        "base_name": "плиты",
        "modifiers": {
            "индук.": "плиты индук.",
            "элек.": "плиты элек.",
            "газ.": "плиты газ."
        }
    },
    "производственные столы": {
        "base_name": "производственные столы",
        "modifiers": {
            "Н": "производственные столы Н",
            "Д": "производственные столы Д"
        }
    },
    "просеиватели": {
        "base_name": "просеиватели",
        "modifiers": {
            "мука": "просеиватели мука",
            "сахар": "просеиватели сахар"
        }
    },
}


# ============================================================
# МОДЕЛИ ДАННЫХ
# ============================================================

@dataclass
class Instruction:
    cleaning_method: str = ""
    instruction_number: str = ""
    product_name: str = ""
    cleaning_technique: str = ""
    concentration: str = ""
    temperature: str = ""
    exposure_time: str = ""
    inventory: str = ""
    frequency: str = ""
    executor: str = ""
    control_method: str = ""

    def to_dict(self) -> dict:
        result = {}
        for key, value in asdict(self).items():
            if value and value.strip():
                result[key] = value
        return result


@dataclass
class ImportObject:
    name: str
    base_name: str
    modifier: Optional[str]
    category: str
    sort_priority: int = 0
    instructions: List[Instruction] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "base_name": self.base_name,
            "modifier": self.modifier,
            "category": self.category,
            "sort_priority": self.sort_priority,
            "instructions": [instr.to_dict() for instr in self.instructions]
        }


# ============================================================
# ПАРСЕР ТЕКСТОВОГО ФАЙЛА
# ============================================================

def parse_objects_from_txt(txt_path: str) -> List[ImportObject]:
    """
    Парсит текстовый файл с объектами.
    Использует MODIFIER_OBJECTS для определения base_name и modifier.
    """
    objects = []
    current_category = None

    with open(txt_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # Заголовок категории
        if line.endswith(':'):
            current_category = line.rstrip(':')
            continue

        if current_category is None:
            continue

        # Обработка строки объекта
        obj = parse_object_line(line, current_category)
        if obj:
            objects.append(obj)

    return objects


def parse_object_line(line: str, category: str) -> Optional[ImportObject]:
    """Парсит строку объекта с учетом MODIFIER_OBJECTS"""
    name = line.strip()

    # Поиск в MODIFIER_OBJECTS
    for base_full_name, info in MODIFIER_OBJECTS.items():
        base_name = info["base_name"]

        # Проверка: это базовый объект?
        if name == base_full_name:
            return ImportObject(
                name=name,
                base_name=base_name,
                modifier=None,
                category=category
            )

        # Проверка: это модификатор?
        for modifier, full_name in info["modifiers"].items():
            if name == full_name:
                return ImportObject(
                    name=name,
                    base_name=base_name,
                    modifier=modifier,
                    category=category
                )

    # Цельный объект (не из списка модификаторов)
    return ImportObject(
        name=name,
        base_name=name,
        modifier=None,
        category=category
    )


# ============================================================
# ПАРСЕР EXCEL-ФАЙЛА
# ============================================================

def parse_instructions_from_excel(xlsx_path: str) -> Dict[str, List[Instruction]]:
    """Парсит Excel и возвращает словарь {название_объекта: [инструкции]}"""
    instructions_map = defaultdict(list)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    col_indices = {
        'object': 0, 'method': 1, 'instruction_no': 2, 'product': 3,
        'technique': 4, 'concentration': 5, 'temperature': 6, 'exposure': 7,
        'inventory': 8, 'frequency': 9, 'executor': 10, 'control': 11,
    }

    for row in range(2, ws.max_row + 1):
        object_name = ws.cell(row=row, column=col_indices['object'] + 1).value
        if not object_name or not str(object_name).strip():
            continue

        object_name = str(object_name).strip()

        if object_name.endswith(':'):
            continue

        instruction = Instruction(
            cleaning_method=str(ws.cell(row=row, column=col_indices['method'] + 1).value or ""),
            instruction_number=str(ws.cell(row=row, column=col_indices['instruction_no'] + 1).value or ""),
            product_name=str(ws.cell(row=row, column=col_indices['product'] + 1).value or ""),
            cleaning_technique=str(ws.cell(row=row, column=col_indices['technique'] + 1).value or ""),
            concentration=str(ws.cell(row=row, column=col_indices['concentration'] + 1).value or ""),
            temperature=str(ws.cell(row=row, column=col_indices['temperature'] + 1).value or ""),
            exposure_time=str(ws.cell(row=row, column=col_indices['exposure'] + 1).value or ""),
            inventory=str(ws.cell(row=row, column=col_indices['inventory'] + 1).value or ""),
            frequency=str(ws.cell(row=row, column=col_indices['frequency'] + 1).value or ""),
            executor=str(ws.cell(row=row, column=col_indices['executor'] + 1).value or ""),
            control_method=str(ws.cell(row=row, column=col_indices['control'] + 1).value or "")
        )

        instructions_map[object_name].append(instruction)

    return instructions_map


# ============================================================
# ГЕНЕРАЦИЯ JSON
# ============================================================

def generate_import_json(
        txt_path: str,
        xlsx_path: str,
        output_path: str,
        category_sort_order: Dict[str, int] = None
):
    print("🔧 ПОДГОТОВКА JSON ДЛЯ ИМПОРТА")
    print("=" * 50)

    # 1. Парсим объекты
    print(f"\n📄 Чтение объектов из: {txt_path}")
    objects = parse_objects_from_txt(txt_path)
    print(f"   Найдено объектов: {len(objects)}")

    # 2. Парсим инструкции
    print(f"\n📊 Чтение инструкций из: {xlsx_path}")
    instructions_map = parse_instructions_from_excel(xlsx_path)
    total_instructions = sum(len(instrs) for instrs in instructions_map.values())
    print(f"   Найдено объектов с инструкциями: {len(instructions_map)}")
    print(f"   Всего инструкций: {total_instructions}")

    # 3. Связываем инструкции с объектами
    print("\n🔗 Связывание инструкций с объектами...")
    matched = 0
    not_matched = []

    for obj in objects:
        if obj.name in instructions_map:
            obj.instructions = instructions_map[obj.name]
            matched += 1
        else:
            not_matched.append(obj.name)

    print(f"   Совпало: {matched}")
    print(f"   Не совпало (будут без инструкций): {len(not_matched)}")

    # 4. Категории
    if category_sort_order is None:
        categories_order = [
            "Поверхности", "Сантехническое оборудование", "Санитарный пост",
            "Мебель", "Офисная техника", "Многоразовые резиновые СИЗ",
            "Бытовая техника", "Инвентарь, посуда и т.д.",
            "Моечный, уборочный инвентарь и оборудование",
            "Посудомоечное оборудование", "Холодильное оборудование",
            "Дозирующее оборудование", "Тепловое оборудование",
            "Технологическое оборудование", "Упаковочное оборудование"
        ]
        category_sort_order = {name: idx for idx, name in enumerate(categories_order, 1)}

    categories = [{"name": name, "sort_order": order} for name, order in category_sort_order.items()]

    # 5. Сохраняем JSON
    output_data = {
        "version": "1.0",
        "created_at": "2026-04-02",
        "categories": categories,
        "objects": [obj.to_dict() for obj in objects]
    }

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)

    print(f"\n✅ JSON файл создан: {output_path}")
    print(f"   Размер: {Path(output_path).stat().st_size / 1024:.1f} KB")

    if not_matched:
        print(f"\n⚠️ ВНИМАНИЕ: Для {len(not_matched)} объектов нет инструкций:")
        for name in not_matched[:20]:
            print(f"   - {name}")
        if len(not_matched) > 20:
            print(f"   ... и еще {len(not_matched) - 20} объектов")

    return output_data


def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--txt", default="полный перечень всех поверхностей из чек-листа.txt")
    parser.add_argument("--xlsx", default="Объекты для базы данных.xlsx")
    parser.add_argument("--output", "-o", default="import_data.json")
    args = parser.parse_args()

    if not Path(args.txt).exists():
        print(f"❌ Файл не найден: {args.txt}")
        return
    if not Path(args.xlsx).exists():
        print(f"❌ Файл не найден: {args.xlsx}")
        return

    generate_import_json(args.txt, args.xlsx, args.output)


if __name__ == "__main__":
    main()