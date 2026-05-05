#!/usr/bin/env python3
"""
Скрипт сопоставления объектов из БД с графиком периодичности (Excel)
и формирования CSV для загрузки в базу данных.

Регистронезависимый поиск и сопоставление.
"""

import sys
import csv
import re
from pathlib import Path
from collections import defaultdict
import openpyxl

sys.path.insert(0, str(Path(__file__).parent))

from db.database import SessionLocal
from db.models import Object, Instruction, Category, RoomCategory


# ============================================================
# НОРМАЛИЗАЦИЯ
# ============================================================
def normalize_name(name: str) -> str:
    """Упрощённая нормализация для сопоставления (регистронезависимая)"""
    if not name:
        return ""
    name = name.lower().strip()
    name = re.sub(r'[^\w\s]', '', name)
    name = re.sub(r'\s+', ' ', name)
    return name.strip()


def split_compound_objects(text: str) -> list:
    """
    Разбивает составную строку на отдельные объекты.
    Пример: "стены (до 2 м), кабель-каналы (до 2 м), розетки, двери"
    → ["стены (до 2 м)", "кабель-каналы (до 2 м)", "розетки", "двери"]
    """
    parts = []
    current = ""
    parens = 0

    for char in text:
        if char == '(':
            parens += 1
            current += char
        elif char == ')':
            parens -= 1
            current += char
        elif char == ',' and parens == 0:
            parts.append(current.strip())
            current = ""
        else:
            current += char

    if current.strip():
        parts.append(current.strip())

    return [p.strip() for p in parts if p.strip()]


# ============================================================
# КАТЕГОРИИ ПОМЕЩЕНИЙ (нижний регистр для сравнения)
# ============================================================
ROOM_CATEGORIES_LOWER = {
    'бытовое': 'Бытовое',
    'производственное': 'Производственное',
    'складское': 'Складское',
    'санитарное': 'Санитарное',
    'вспомогательное': 'Вспомогательное',
    'моечное': 'Моечное',
    'техническое': 'Техническое',
    'офисное': 'Офисное',
    'общего назначения': 'Общего назначения',
}


def get_room_category(text: str) -> str:
    """Определяет категорию помещения по тексту (регистронезависимо)"""
    text_lower = text.lower().strip()

    if text_lower in ROOM_CATEGORIES_LOWER:
        return ROOM_CATEGORIES_LOWER[text_lower]

    for key, value in ROOM_CATEGORIES_LOWER.items():
        if key in text_lower or text_lower in key:
            return value

    return None


# ============================================================
# ЧТЕНИЕ EXCEL
# ============================================================
def read_excel(file_path: str) -> dict:
    """
    Читает Excel с графиком периодичности.
    Возвращает:
    {
        "Бытовое": {
            "объект": [
                {"maintenance_type": "основная", "frequency_wash": "...", "frequency_des": "...", "executor": "..."},
                ...
            ],
            ...
        },
        ...
    }
    """
    file_path = Path(file_path)
    if not file_path.exists():
        print(f"❌ Файл не найден: {file_path}")
        return {}

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    ws = wb.active

    print(f"   Лист: {ws.title}, Строк: {ws.max_row}, Колонок: {ws.max_column}")

    result = defaultdict(lambda: defaultdict(list))
    current_category = None
    current_object = None
    header_skipped = False

    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row:
            continue

        cells = [str(c).strip() if c is not None else '' for c in row]
        while len(cells) < 6:
            cells.append('')

        a, b, c, d, e, f = cells[0], cells[1], cells[2], cells[3], cells[4], cells[5]

        # Пропускаем строку заголовка
        if not header_skipped and a == 'Описание работ/объект обработки':
            header_skipped = True
            continue

        b_lower = b.lower()

        # Ищем категорию помещения в колонке F
        room_cat = get_room_category(f)
        if room_cat:
            current_category = room_cat
            if a:
                current_object = a

        if not current_category:
            # Пробуем найти категорию в A (для строк-разделителей)
            room_cat = get_room_category(a)
            if room_cat:
                current_category = room_cat
            continue

        # Пропускаем заголовки разделов
        if a and not b:
            a_lower = a.lower()
            if any(s in a_lower for s in ['удаление отходов', 'санитарная обработка']):
                continue
            current_object = a
            continue

        # Строка с уровнем обработки
        if b_lower in ['основная', 'поддерживающая', 'генеральная']:
            if a:
                current_object = a

            if current_object:
                result[current_category][current_object].append({
                    'maintenance_type': b,
                    'frequency_wash': c if c else None,
                    'frequency_des': d if d else None,
                    'executor': e if e else None,
                })

    return result


# ============================================================
# ПОИСК ОБЪЕКТА В БД
# ============================================================
def find_object_in_db(session, obj_name: str) -> list:
    """
    Ищет объекты в БД по имени (составные строки разбираются).
    Возвращает список объектов.
    """
    found = []

    # Разбиваем составную строку
    sub_names = split_compound_objects(obj_name)

    all_objects = session.query(Object).all()

    for sub_name in sub_names:
        sub_norm = normalize_name(sub_name)

        for obj in all_objects:
            obj_norm = normalize_name(obj.normalized_name or '')
            obj_disp = normalize_name(obj.display_name or '')
            obj_base = normalize_name(obj.base_name or '')

            if (sub_norm == obj_norm or
                    sub_norm == obj_disp or
                    sub_norm == obj_base or
                    sub_norm in obj_norm or obj_norm in sub_norm or
                    sub_norm in obj_disp or obj_disp in sub_norm or
                    sub_norm in obj_base or obj_base in sub_norm):
                if obj not in found:
                    found.append(obj)

    return found


# ============================================================
# ФОРМИРОВАНИЕ CSV
# ============================================================
def generate_csv(excel_data: dict, output_path: str):
    """
    Сопоставляет данные из Excel с БД и формирует CSV.
    """
    session = SessionLocal()

    room_cats = {rc.name: rc.id for rc in session.query(RoomCategory).all()}
    obj_cats = {c.id: c.name for c in session.query(Category).all()}
    all_objects = session.query(Object).order_by(Object.category_id, Object.display_name).all()

    rows = []

    for room_cat_name, objects_data in excel_data.items():
        print(f"\n🏢 Категория помещения: {room_cat_name} ({len(objects_data)} объектов)")

        matched_objects = set()

        for excel_obj_name, levels in objects_data.items():
            db_objects = find_object_in_db(session, excel_obj_name)

            if not db_objects:
                print(f"  ⚠️ Не найден: {excel_obj_name}")
                continue

            for db_obj in db_objects:
                matched_objects.add(db_obj.id)

                for level in levels:
                    maintenance_type = level['maintenance_type']
                    executor = level['executor']

                    if level['frequency_wash']:
                        rows.append({
                            'object_id': db_obj.id,
                            'category_name': obj_cats.get(db_obj.category_id, ''),
                            'normalized_name': db_obj.normalized_name,
                            'display_name': db_obj.display_name,
                            'base_name': db_obj.base_name,
                            'modifier': db_obj.modifier or '',
                            'sort_priority': db_obj.sort_priority or '',
                            'instruction_id': '',
                            'room_category_name': room_cat_name,
                            'maintenance_type': maintenance_type,
                            'cleaning_method': 'мойка',
                            'instruction_number': '',
                            'product_name': '',
                            'cleaning_technique': '',
                            'concentration': '',
                            'temperature': '',
                            'exposure_time': '',
                            'inventory': '',
                            'frequency': level['frequency_wash'],
                            'executor': executor,
                            'control_method': '',
                        })

                    if level['frequency_des']:
                        rows.append({
                            'object_id': db_obj.id,
                            'category_name': obj_cats.get(db_obj.category_id, ''),
                            'normalized_name': db_obj.normalized_name,
                            'display_name': db_obj.display_name,
                            'base_name': db_obj.base_name,
                            'modifier': db_obj.modifier or '',
                            'sort_priority': db_obj.sort_priority or '',
                            'instruction_id': '',
                            'room_category_name': room_cat_name,
                            'maintenance_type': maintenance_type,
                            'cleaning_method': 'дезинфекция',
                            'instruction_number': '',
                            'product_name': '',
                            'cleaning_technique': '',
                            'concentration': '',
                            'temperature': '',
                            'exposure_time': '',
                            'inventory': '',
                            'frequency': level['frequency_des'],
                            'executor': executor,
                            'control_method': '',
                        })

        # Несопоставленные объекты БД
        unmatched = 0
        for obj in all_objects:
            if obj.id not in matched_objects:
                rows.append({
                    'object_id': obj.id,
                    'category_name': obj_cats.get(obj.category_id, ''),
                    'normalized_name': obj.normalized_name,
                    'display_name': obj.display_name,
                    'base_name': obj.base_name,
                    'modifier': obj.modifier or '',
                    'sort_priority': obj.sort_priority or '',
                    'instruction_id': '',
                    'room_category_name': room_cat_name,
                    'maintenance_type': 'основная',
                    'cleaning_method': 'мойка',
                    'instruction_number': '',
                    'product_name': '',
                    'cleaning_technique': '',
                    'concentration': '',
                    'temperature': '',
                    'exposure_time': '',
                    'inventory': '',
                    'frequency': '',
                    'executor': '',
                    'control_method': '',
                })
                rows.append({
                    'object_id': obj.id,
                    'category_name': obj_cats.get(obj.category_id, ''),
                    'normalized_name': obj.normalized_name,
                    'display_name': obj.display_name,
                    'base_name': obj.base_name,
                    'modifier': obj.modifier or '',
                    'sort_priority': obj.sort_priority or '',
                    'instruction_id': '',
                    'room_category_name': room_cat_name,
                    'maintenance_type': 'основная',
                    'cleaning_method': 'дезинфекция',
                    'instruction_number': '',
                    'product_name': '',
                    'cleaning_technique': '',
                    'concentration': '',
                    'temperature': '',
                    'exposure_time': '',
                    'inventory': '',
                    'frequency': '',
                    'executor': '',
                    'control_method': '',
                })
                unmatched += 1

        print(f"  ✅ Найдено: {len(matched_objects)}")
        print(f"  ➕ Без инструкций: {unmatched}")

    session.close()

    # Сохраняем CSV
    fieldnames = [
        'object_id', 'category_name', 'normalized_name', 'display_name', 'base_name',
        'modifier', 'sort_priority', 'instruction_id', 'room_category_name',
        'maintenance_type', 'cleaning_method', 'instruction_number', 'product_name',
        'cleaning_technique', 'concentration', 'temperature', 'exposure_time',
        'inventory', 'frequency', 'executor', 'control_method'
    ]

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=';')
        writer.writeheader()
        writer.writerows(rows)

    print(f"\n{'=' * 60}")
    print(f"📁 Сохранено: {output_path}")
    print(f"📋 Всего строк: {len(rows)}")


# ============================================================
# MAIN
# ============================================================
def main():
    if len(sys.argv) < 2:
        print("Использование: python match_schedule.py <Excel_файл> [выходной_CSV]")
        print("Пример: python match_schedule.py 'периодичность обработки всех помещений.xlsx'")
        return

    excel_file = sys.argv[1]
    output_csv = sys.argv[2] if len(sys.argv) > 2 else "output.csv"

    print("=" * 60)
    print("🔍 СОПОСТАВЛЕНИЕ EXCEL С БД")
    print("=" * 60)

    data = read_excel(excel_file)

    if not data:
        print("❌ Данные не прочитаны. Проверьте файл Excel.")
        return

    print(f"\n📊 Категорий помещений в Excel: {len(data)}")
    for cat in data:
        print(f"  - {cat}: {len(data[cat])} объектов")

    print(f"\n🔍 Сопоставление с БД...")
    generate_csv(data, output_csv)


if __name__ == "__main__":
    main()