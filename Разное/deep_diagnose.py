from docx import Document
from pathlib import Path


def deep_diagnose(template_path="tech_card_templates/шаблон.docx"):
    doc = Document(template_path)

    if not doc.tables:
        print("❌ Нет таблиц")
        return

    table = doc.tables[0]
    print(f"📊 Таблица: {len(table.rows)} строк, {len(table.columns)} колонок (по атрибутам)")

    # Детальный разбор каждой строки и ячейки
    for row_idx, row in enumerate(table.rows):
        print(f"\n--- Строка {row_idx} ---")
        print(f"  Количество ячеек в строке (len(row.cells)): {len(row.cells)}")
        for col_idx, cell in enumerate(row.cells):
            # Получаем текст через внутренний XML
            text = cell.text
            # Проверяем наличие подстроки "Объект"
            has_object = "Объект" in text
            # Выводим информацию
            print(
                f"    Ячейка [{row_idx},{col_idx}]: текст='{text[:50] if text else '(пусто)'}', содержит 'Объект': {has_object}")
            # Если текст пустой, но возможно есть runs
            if not text and cell.paragraphs:
                for p_idx, para in enumerate(cell.paragraphs):
                    for r_idx, run in enumerate(para.runs):
                        print(
                            f"        Run [{p_idx},{r_idx}]: text='{run.text}', font={run.font.name}, size={run.font.size}")

    # Ищем строку, где есть слово "Объект" в любой ячейке
    print("\n" + "=" * 50)
    print("🔎 Поиск 'Объект обработки'...")
    found = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            if "Объект" in cell.text:
                print(f"✅ Найдено в строке {row_idx}, колонке {col_idx}: '{cell.text[:100]}'")
                found = True
    if not found:
        print("❌ Слово 'Объект' не найдено ни в одной ячейке.")
        print("   Возможно, шаблон поврежден или имеет другую структуру.")


if __name__ == "__main__":
    deep_diagnose()