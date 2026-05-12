# check_row7.py
from docx import Document
from docx.oxml.ns import qn

doc = Document("tech_card_templates/шаблон.docx")
table = doc.tables[1]

print(f"Всего строк: {len(table.rows)}")

row = table.rows[7]
for col_idx, cell in enumerate(row.cells):
    print(f"\nКолонка {col_idx}:")
    for p_idx, para in enumerate(cell.paragraphs):
        print(f"  Абзац {p_idx}: text='{para.text[:40]}', runs={len(para.runs)}")
        for r_idx, run in enumerate(para.runs):
            r_pr = run._r.find(qn('w:rPr'))
            if r_pr is not None:
                bold = r_pr.find(qn('w:b'))
                sz = r_pr.find(qn('w:sz'))
                print(f"    Run {r_idx}: text='{run.text[:30]}', bold={'✅' if bold is not None else '❌'}, sz={'✅' if sz is not None else '❌'}")
            else:
                print(f"    Run {r_idx}: text='{run.text[:30]}', rPr=❌")