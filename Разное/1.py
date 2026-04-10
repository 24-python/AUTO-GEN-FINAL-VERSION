# create_correct_template.py
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path


def create_template():
    doc = Document()

    # Настройка полей
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(0.5)

    # Создаем таблицу: 7 строк, 12 колонок
    table = doc.add_table(rows=7, cols=12)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # ---- СТРОКА 0: заголовок (объединенная) ----
    cell = table.cell(0, 0)
    for col in range(1, 12):
        cell.merge(table.cell(0, col))
    cell.text = "Технологическая карта санитарной обработки № ________"
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].runs[0]
    run.font.size = Pt(14)
    run.font.bold = True

    # ---- СТРОКА 1: Помещение (объединенная) ----
    cell = table.cell(1, 0)
    for col in range(1, 12):
        cell.merge(table.cell(1, col))
    cell.text = "Помещение: _________________________"

    # ---- СТРОКА 2: пустая (объединенная) ----
    cell = table.cell(2, 0)
    for col in range(1, 12):
        cell.merge(table.cell(2, col))
    cell.text = ""

    # ---- СТРОКА 3: примечание (объединенная) ----
    cell = table.cell(3, 0)
    for col in range(1, 12):
        cell.merge(table.cell(3, col))
    cell.text = "Санитарная обработка помещений осуществляется после полного удаления из них сырьевых компонентов и материалов."
    run = cell.paragraphs[0].runs[0]
    run.font.size = Pt(10)
    run.italic = True

    # ---- СТРОКА 4: пустая (объединенная) ----
    cell = table.cell(4, 0)
    for col in range(1, 12):
        cell.merge(table.cell(4, col))
    cell.text = ""

    # ---- СТРОКА 5: ЗАГОЛОВКИ КОЛОНОК ----
    headers = [
        "Объект обработки", "Способ обработки", "№ инструкции",
        "Наименование средства", "Метод уборки",
        "Концентрация средства, %, метод разведения",
        "Температура раствора, °C", "Время выдержки средства, мин",
        "Инвентарь", "Периодичность", "Исполнитель",
        "Метод контроля/периодичность"
    ]
    for col, header in enumerate(headers):
        cell = table.cell(5, col)
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(9)

    # ---- СТРОКА 6: пустая (образец для данных) ----
    for col in range(12):
        table.cell(6, col).text = ""

    # Настройка ширины колонок
    widths = [Cm(3.0), Cm(2.0), Cm(1.5), Cm(2.5), Cm(2.0), Cm(2.5),
              Cm(1.8), Cm(1.8), Cm(1.8), Cm(2.0), Cm(2.0), Cm(2.5)]
    for col, width in enumerate(widths):
        table.columns[col].width = width

    # Сохраняем
    output_path = Path("tech_card_templates/шаблон.docx")
    output_path.parent.mkdir(exist_ok=True)
    doc.save(str(output_path))
    print(f"✅ Шаблон создан: {output_path}")
    print("   Теперь это настоящая таблица Word, а не ASCII-графика")


if __name__ == "__main__":
    # Удаляем старый файл, если есть
    import os

    old_file = Path("tech_card_templates/шаблон.docx")
    if old_file.exists():
        os.remove(old_file)
        print(f"🗑️ Удален старый файл")
    create_template()