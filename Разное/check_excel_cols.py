# check_excel_cols.py
import openpyxl

wb = openpyxl.load_workbook("периодичность обработки всех помещений.xlsx", data_only=True)
ws = wb.active

print(f"Строк: {ws.max_row}, Колонок: {ws.max_column}")

for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
    cells = [str(c)[:60] if c is not None else 'None' for c in row]
    print(cells)