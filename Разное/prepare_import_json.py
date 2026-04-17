#!/usr/bin/env python3
"""
prepare_import_json.py

Подготавливает JSON для импорта в БД на основе:
1. CSV-файла с соответствиями (normalized_name → display_name → category)
2. Excel-файла с инструкциями (Объекты для базы данных.xlsx)

Создаёт import_data.json со структурой:
{
    "categories": [...],
    "objects": [
        {
            "normalized_name": "холод столы",
            "display_name": "холодильные столы",
            "base_name": "холодильные столы",
            "modifier": null,
            "category": "Холодильное оборудование",
            "sort_priority": 0,
            "instructions": [...]
        }
    ]
}
"""

import csv
import json
import re
from pathlib import Path
from collections import defaultdict
from typing import Dict, List, Optional

import openpyxl


# ============================================================
# ФУНКЦИЯ НОРМАЛИЗАЦИИ (ДОЛЖНА СОВПАДАТЬ С ПАРСЕРОМ)
# ============================================================

def normalize_name(name: str) -> str:
    """
    Нормализует имя объекта так же, как это делает парсер.
    Удаляет знаки препинания, оставляет буквы, цифры, пробелы, дефисы, скобки.
    """
    if not name:
        return ""
    # Удаляем всё, кроме букв, цифр, пробелов, дефисов, скобок
    normalized = re.sub(r'[^\w\s\-\(\)]', '', name)
    # Приводим к нижнему регистру
    normalized = normalized.lower()
    # Заменяем множественные пробелы на один
    normalized = re.sub(r'\s+', ' ', normalized)
    return normalized.strip()


# ============================================================
# ЗАГРУЗКА СООТВЕТСТВИЙ ИЗ CSV
# ============================================================

def load_normalization_map(csv_path: str) -> Dict[str, dict]:
    """
    Загружает CSV с колонками: normalized_name;display_name;category
    Возвращает словарь: {normalized_name: {"display_name": ..., "category": ...}}
    """
    mapping = {}
    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f, delimiter=';')
        for row in reader:
            if len(row) < 3:
                continue
            normalized = row[0].strip()
            display_name = row[1].strip()
            category = row[2].strip()
            if normalized and display_name and category:
                mapping[normalized] = {
                    "display_name": display_name,
                    "category": category
                }
    return mapping


# ============================================================
# ЗАГРУЗКА КАТЕГОРИЙ ИЗ CSV
# ============================================================

def load_categories_from_csv(csv_path: str) -> List[str]:
    """Извлекает уникальные категории из CSV."""
    categories = set()
    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f, delimiter=';')
        for row in reader:
            if len(row) >= 3:
                categories.add(row[2].strip())
    return sorted(categories)


# ============================================================
# ПАРСЕР ИНСТРУКЦИЙ ИЗ EXCEL
# ============================================================

def parse_instructions_from_excel(xlsx_path: str) -> Dict[str, List[dict]]:
    """
    Парсит Excel с инструкциями.
    Возвращает словарь: {normalized_name: [список_инструкций]}
    """
    instructions_map = defaultdict(list)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    col_indices = {
        'object': 0, 'method': 1, 'instruction_no': 2, 'product': 3,
        'technique': 4, 'concentration': 5, 'temperature': 6, 'exposure': 7,
        'inventory': 8, 'frequency': 9, 'executor': 10, 'control': 11,
    }

    for row in range(2, ws.max_row + 1):
        object_name = ws.cell(row=row, column=col_indices['object'] + 1).value
        if not object_name or not str(object_name).strip():
            continue

        object_name = str(object_name).strip()

        # Пропускаем строки-заголовки категорий
        if object_name.endswith(':'):
            continue

        # В Excel теперь уже нормализованные имена, используем их как есть
        normalized = object_name

        instruction = {
            "cleaning_method": str(ws.cell(row=row, column=col_indices['method'] + 1).value or ""),
            "instruction_number": str(ws.cell(row=row, column=col_indices['instruction_no'] + 1).value or ""),
            "product_name": str(ws.cell(row=row, column=col_indices['product'] + 1).value or ""),
            "cleaning_technique": str(ws.cell(row=row, column=col_indices['technique'] + 1).value or ""),
            "concentration": str(ws.cell(row=row, column=col_indices['concentration'] + 1).value or ""),
            "temperature": str(ws.cell(row=row, column=col_indices['temperature'] + 1).value or ""),
            "exposure_time": str(ws.cell(row=row, column=col_indices['exposure'] + 1).value or ""),
            "inventory": str(ws.cell(row=row, column=col_indices['inventory'] + 1).value or ""),
            "frequency": str(ws.cell(row=row, column=col_indices['frequency'] + 1).value or ""),
            "executor": str(ws.cell(row=row, column=col_indices['executor'] + 1).value or ""),
            "control_method": str(ws.cell(row=row, column=col_indices['control'] + 1).value or ""),
        }

        instructions_map[normalized].append(instruction)

    return instructions_map


# ============================================================
# ГЕНЕРАЦИЯ JSON
# ============================================================

def generate_import_json(csv_path: str, xlsx_path: str, output_path: str):
    """Генерирует JSON для импорта в БД."""

    print("🔧 ПОДГОТОВКА JSON ДЛЯ ИМПОРТА")
    print("=" * 50)

    # 1. Загружаем соответствия из CSV
    print(f"\n📄 Чтение соответствий из: {csv_path}")
    mapping = load_normalization_map(csv_path)
    print(f"   Загружено соответствий: {len(mapping)}")

    # 2. Загружаем категории
    categories_list = load_categories_from_csv(csv_path)
    categories = [
        {"name": name, "sort_order": idx + 1}
        for idx, name in enumerate(categories_list)
    ]
    print(f"   Категорий: {len(categories)}")

    # 3. Парсим инструкции из Excel
    print(f"\n📊 Чтение инструкций из: {xlsx_path}")
    instructions_map = parse_instructions_from_excel(xlsx_path)
    total_instructions = sum(len(instrs) for instrs in instructions_map.values())
    print(f"   Уникальных объектов с инструкциями: {len(instructions_map)}")
    print(f"   Всего инструкций: {total_instructions}")

    # 4. Формируем объекты для JSON
    print("\n🔗 Формирование объектов...")
    objects = []
    matched = 0
    not_matched = []

    for normalized, data in mapping.items():
        display_name = data["display_name"]
        category = data["category"]

        # Получаем инструкции для этого normalized_name
        instructions = instructions_map.get(normalized, [])
        if instructions:
            matched += 1
        else:
            not_matched.append(normalized)

        # Определяем base_name и modifier (упрощённо)
        base_name = display_name
        modifier = None

        # Пытаемся извлечь модификатор для отображения
        if "П" in normalized or "ОК" in normalized or "Н" in normalized or "А" in normalized or "С" in normalized:
            parts = normalized.split()
            if len(parts) > 1 and len(parts[-1]) <= 3:
                modifier = parts[-1]
                base_name = " ".join(parts[:-1])

        obj = {
            "normalized_name": normalized,
            "display_name": display_name,
            "base_name": base_name,
            "modifier": modifier,
            "category": category,
            "sort_priority": 0,
            "instructions": instructions
        }
        objects.append(obj)

    print(f"   Объектов с инструкциями: {matched}")
    print(f"   Объектов без инструкций: {len(not_matched)}")

    # 5. Сохраняем JSON
    output_data = {
        "version": "1.0",
        "created_at": "2026-04-09",
        "categories": categories,
        "objects": objects
    }

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)

    print(f"\n✅ JSON файл создан: {output_path}")
    print(f"   Размер: {Path(output_path).stat().st_size / 1024:.1f} KB")

    if not_matched:
        print(f"\n⚠️ ВНИМАНИЕ: Для {len(not_matched)} объектов нет инструкций в Excel:")
        for name in not_matched[:20]:
            print(f"   - {name}")
        if len(not_matched) > 20:
            print(f"   ... и ещё {len(not_matched) - 20}")
    else:
        print("\n✅ Все объекты имеют инструкции!")


# ============================================================
# ТОЧКА ВХОДА
# ============================================================

def main():
    import argparse

    parser = argparse.ArgumentParser(description="Подготовка JSON для импорта в БД")
    parser.add_argument("--csv", default="сопостовление объектов для БД.csv",
                        help="Путь к CSV с соответствиями normalized_name;display_name;category")
    parser.add_argument("--xlsx", default="Объекты для базы данных.xlsx",
                        help="Путь к Excel-файлу с инструкциями")
    parser.add_argument("--output", "-o", default="import_data.json",
                        help="Выходной JSON файл")
    args = parser.parse_args()

    if not Path(args.csv).exists():
        print(f"❌ CSV файл не найден: {args.csv}")
        return
    if not Path(args.xlsx).exists():
        print(f"❌ Excel файл не найден: {args.xlsx}")
        return

    generate_import_json(args.csv, args.xlsx, args.output)


if __name__ == "__main__":
    main()