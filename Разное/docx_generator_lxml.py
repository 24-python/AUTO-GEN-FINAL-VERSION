"""
Быстрый генератор техкарт с использованием lxml для заполнения ячеек.
"""

from docx import Document
from pathlib import Path
from parser.models import ChecklistData
from db.models import Instruction, Category as DBCategory, Object as DBObject
from db.database import SessionLocal
from collections import defaultdict
import zipfile
import tempfile
import time
import re
from lxml import etree


class FastTechCardGenerator:
    TEMPLATES_DIR = Path("tech_card_templates")
    DEFAULT_TEMPLATE = TEMPLATES_DIR / "шаблон.docx"

    PRODUCT_COLORS = {
        "ХИМИТЕК ПОЛИДЕЗ®-СУПЕР": "FFFFCC",
        "ХИМИТЕК УНИВЕРСАЛ-ПД-Н": "99FF99",
        "ХИМИТЕК ЧУДОДЕЙ®-КОМБИ-ПЕНАКТИВ": "99CCFF",
        "ХИМИТЕК ЧУДОДЕЙ®-CIP": "99CCFF",
        "ХИМИТЕК ПОЛИКОР®": "FFCCCC",
        "ХИМИТЕК ЧАРОЙТ®-СПРЕЙ": "99FF99",
        "ХИМИТЕК СВЕЖЕСТЬ-АНТИСЕПТИК": "FFFFCC",
    }

    NAMESPACES = {
        'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
    }

    def __init__(self, template_path: str = None):
        self.template_path = Path(template_path) if template_path else self.DEFAULT_TEMPLATE

    def _normalize_product_name(self, name: str) -> str:
        if not name:
            return ""
        normalized = re.sub(r'[^\w\s\-]', '', name)
        normalized = normalized.lower()
        normalized = re.sub(r'\s+', ' ', normalized)
        return normalized.strip()

    def _get_color_for_product(self, product_name: str) -> str:
        if not product_name:
            return None
        normalized_input = self._normalize_product_name(product_name)
        for key, color in self.PRODUCT_COLORS.items():
            if self._normalize_product_name(key) == normalized_input:
                return color
        return None

    def _get_category_priority(self, session) -> dict:
        categories = session.query(DBCategory).order_by(DBCategory.sort_order).all()
        return {cat.name: cat.sort_order for cat in categories}

    def _set_cell_text(self, cell, text: str):
        """Устанавливает текст в ячейке"""
        t_elements = cell.xpath('.//w:t', namespaces=self.NAMESPACES)

        if t_elements:
            t_elements[0].text = text
        else:
            p = cell.find('.//w:p', namespaces=self.NAMESPACES)
            if p is None:
                p = etree.SubElement(cell, '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')

            r = etree.SubElement(p, '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r')
            t = etree.SubElement(r, '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
            t.text = text

    def _clean_row_attributes(self, row_element):
        """Удаляет служебные атрибуты для уменьшения размера файла"""
        ATTRS_TO_REMOVE = [
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rsidR',
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rsidRPr',
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rsidTr',
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rsidRDefault',
            '{http://schemas.microsoft.com/office/word/2010/wordml}paraId',
            '{http://schemas.microsoft.com/office/word/2010/wordml}textId',
        ]

        for elem in row_element.iter():
            for attr in ATTRS_TO_REMOVE:
                if attr in elem.attrib:
                    del elem.attrib[attr]

    def _prepare_rows_data(self, checklist_data: ChecklistData) -> tuple:
        """Подготавливает данные для заполнения таблицы"""
        session = SessionLocal()
        category_priority = self._get_category_priority(session)

        checked_items = checklist_data.get_checked_items()
        all_names = [item.name for item in checked_items]

        db_objects = session.query(DBObject).filter(DBObject.normalized_name.in_(all_names)).all()
        objects_dict = {obj.normalized_name: obj for obj in db_objects}

        object_ids = [obj.id for obj in db_objects]
        db_instructions = session.query(Instruction).filter(Instruction.object_id.in_(object_ids)).all()

        instructions_dict = defaultdict(list)
        for instr in db_instructions:
            instructions_dict[instr.object_id].append(instr)

        category_items = defaultdict(list)
        category_order = {}

        for item in checked_items:
            obj = objects_dict.get(item.name)
            cat_name = item.category.value
            priority = category_priority.get(cat_name, 999)
            category_order[cat_name] = priority

            display_name = obj.display_name if obj else item.name

            if obj and obj.id in instructions_dict:
                for instr in instructions_dict[obj.id]:
                    category_items[cat_name].append((display_name, instr))
            else:
                category_items[cat_name].append((display_name, None))

        session.close()

        sorted_categories = sorted(category_items.keys(), key=lambda x: category_order.get(x, 999))
        for cat_name in sorted_categories:
            category_items[cat_name].sort(key=lambda x: x[0])

        rows_data = []
        for category_name in sorted_categories:
            items = category_items[category_name]
            if not items:
                continue
            rows_data.append(('category', category_name, None))
            for obj_name, instr in items:
                rows_data.append(('object', obj_name, instr))

        return rows_data, len(db_objects), len(db_instructions)

    def _fill_cells_lxml(self, row_element, row_data: tuple):
        """Заполняет ячейки строки через lxml"""
        cells = row_element.xpath('.//w:tc', namespaces=self.NAMESPACES)

        if row_data[0] == 'category':
            if len(cells) >= 1:
                self._set_cell_text(cells[0], row_data[1])

        elif row_data[0] == 'object':
            obj_name, instr = row_data[1], row_data[2]

            if len(cells) >= 12:
                self._set_cell_text(cells[0], obj_name)

                if instr:
                    self._set_cell_text(cells[1], instr.cleaning_method or "")
                    self._set_cell_text(cells[2], instr.instruction_number or "")
                    self._set_cell_text(cells[3], instr.product_name or "")
                    self._set_cell_text(cells[4], instr.cleaning_technique or "")
                    self._set_cell_text(cells[5], instr.concentration or "")
                    self._set_cell_text(cells[6], instr.temperature or "")
                    self._set_cell_text(cells[7], instr.exposure_time or "")
                    self._set_cell_text(cells[8], instr.inventory or "")
                    self._set_cell_text(cells[9], instr.frequency or "")
                    self._set_cell_text(cells[10], instr.executor or "")
                    self._set_cell_text(cells[11], instr.control_method or "")

    def generate(self, checklist_data: ChecklistData, output_path: str) -> str:
        print("      ⚡ lxml-генерация...")
        total_start = time.time()

        # 1. Подготовка данных
        print("      📋 Подготовка данных...")
        t = time.time()
        rows_data, obj_count, instr_count = self._prepare_rows_data(checklist_data)
        print(f"         {time.time() - t:.2f}с (строк: {len(rows_data)})")

        # 2. Создаём временный документ
        print("      📂 Загрузка шаблона...")
        t = time.time()
        doc = Document(self.template_path)

        main_table = doc.tables[1]
        room_cell = main_table.cell(1, 0)
        if "Помещение:" in room_cell.text:
            room_cell.text = f"Помещение: {checklist_data.room_name or '______________'}"

        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            temp_path = tmp.name
        doc.save(temp_path)
        print(f"         {time.time() - t:.2f}с")

        # 3. Работаем с XML через lxml
        print("      ⚡ Заполнение через lxml...")
        t = time.time()

        with zipfile.ZipFile(temp_path, 'a') as zf:
            xml_content = zf.read('word/document.xml')
            root = etree.fromstring(xml_content)

            tables = root.xpath('//w:tbl', namespaces=self.NAMESPACES)
            if len(tables) < 2:
                raise ValueError("Не найдена вторая таблица в документе")
            table = tables[1]

            tr_elements = table.xpath('.//w:tr', namespaces=self.NAMESPACES)

            # Оставляем первые 7 строк (0-6), удаляем остальные
            start_row = 7
            for tr in tr_elements[start_row:]:
                table.remove(tr)

            # Шаблон строки - строка 6
            template_row = tr_elements[6] if len(tr_elements) > 6 else tr_elements[-1]

            # Клонируем и заполняем строки
            for row_data in rows_data:
                new_row = etree.fromstring(etree.tostring(template_row))
                self._clean_row_attributes(new_row)  # ← ОЧИЩАЕМ АТРИБУТЫ
                self._fill_cells_lxml(new_row, row_data)
                table.append(new_row)

            zf.writestr('word/document.xml', etree.tostring(root, encoding='UTF-8', xml_declaration=True))

        print(f"         {time.time() - t:.2f}с")

        # 4. Перемещаем в выходной файл
        print("      💾 Сохранение...")
        t = time.time()
        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        if output_file.exists():
            output_file.unlink()
        Path(temp_path).rename(output_file)
        print(f"         {time.time() - t:.2f}с")

        # Размер файла
        size_kb = output_file.stat().st_size / 1024
        print(f"      📁 Размер файла: {size_kb:.1f} KB")

        print(f"      ✅ ГОТОВО! Общее время: {time.time() - total_start:.2f}с")
        return str(output_file)