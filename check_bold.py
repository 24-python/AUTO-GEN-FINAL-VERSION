# check_all_rows.py
from docx import Document
from docx.oxml.ns import qn

doc = Document("tech_card_templates/шаблон.docx")
table = doc.tables[1]

print(f"Всего строк в таблице: {len(table.rows)}\n")

for row_idx, row in enumerate(table.rows):
    col3 = row.cells[3] if len(row.cells) > 3 else None
    if col3 and col3.text.strip():
        print(f"Строка {row_idx} (индекс), колонка 4: text='{col3.text.strip()[:60]}'")
        for r_idx, run in enumerate(col3.paragraphs[0].runs):
            r_pr = run._r.find(qn('w:rPr'))
            bold = r_pr.find(qn('w:b')) if r_pr is not None else None
            print(f"  Run {r_idx}: bold={'✅' if bold is not None else '❌'}")
        print()